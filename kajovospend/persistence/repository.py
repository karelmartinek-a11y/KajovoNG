from __future__ import annotations

import hashlib
import json
import re
import shutil
import sqlite3
import tempfile
import uuid
from contextlib import contextmanager
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any, Iterator

from kajovospend.app.constants import PROJECT_MARKER
from kajovospend.ocr import DocumentLoader, LoadedDocument
from kajovospend.persistence.project_context import ProjectDatabasePaths, WorkingConnectionFactory, ProductionConnectionFactory
from kajovospend.persistence.repository_reporting_views import apply_reporting_view_bindings


def utc_now() -> str:
    return datetime.now(UTC).replace(microsecond=0).isoformat()


def _parse_timestamp(value: Any) -> datetime | None:
    text = str(value or '').strip()
    if not text:
        return None
    try:
        parsed = datetime.fromisoformat(text.replace('Z', '+00:00'))
    except ValueError:
        return None
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=UTC)
    return parsed.astimezone(UTC)


class ProjectRepository:
    def __init__(self, project_path: Path):
        self.project_path = Path(project_path).expanduser().resolve()
        self.paths = ProjectDatabasePaths.from_project_root(self.project_path)
        self.working_factory = WorkingConnectionFactory(self.paths)
        self.production_factory = ProductionConnectionFactory(self.paths)
        self.document_loader = DocumentLoader()
        self.runtime_log_path = self.project_path / 'logs' / 'runtime.log'
        self.structured_log_path = self.project_path / 'logs' / 'decisions.jsonl'

    @property
    def work_db_path(self) -> Path:
        return self.paths.working_db

    @property
    def prod_db_path(self) -> Path:
        return self.paths.production_db

    @contextmanager
    def work_connection(self) -> Iterator[sqlite3.Connection]:
        conn = self.working_factory.connect()
        try:
            yield conn
        finally:
            conn.close()

    @contextmanager
    def prod_connection(self) -> Iterator[sqlite3.Connection]:
        conn = self.production_factory.connect()
        try:
            yield conn
        finally:
            conn.close()

    def ensure_schema(self) -> None:
        with self.work_connection() as conn:
            conn.executescript(
                '''
                create virtual table if not exists working_documents_fts using fts5(
                    original_name,
                    document_number,
                    supplier_ico,
                    supplier_name,
                    quarantine_reason,
                    last_error,
                    correlation_id
                );
                create virtual table if not exists extraction_attempts_fts using fts5(
                    original_name,
                    reason,
                    branch,
                    processor_name,
                    correlation_id
                );
                create table if not exists operational_log (
                    id integer primary key,
                    created_at text default current_timestamp,
                    level text not null,
                    event_type text not null,
                    message text not null,
                    correlation_id text,
                    file_id integer,
                    document_id integer,
                    attempt_id integer,
                    payload_json text
                );
                create table if not exists errors (
                    id integer primary key,
                    created_at text default current_timestamp,
                    correlation_id text,
                    file_id integer,
                    document_id integer,
                    attempt_id integer,
                    error_code text,
                    error_message text,
                    severity text default 'error',
                    payload_json text
                );
                create table if not exists document_results (
                    id integer primary key,
                    file_id integer not null,
                    document_id integer,
                    final_state text not null,
                    reason text,
                    decided_at text default current_timestamp,
                    correlation_id text,
                    final_document_id integer,
                    unique(file_id)
                );
                create table if not exists audit_changes (
                    id integer primary key,
                    created_at text default current_timestamp,
                    correlation_id text,
                    entity_type text not null,
                    entity_id integer,
                    action text not null,
                    before_json text,
                    after_json text,
                    note text
                );
                create table if not exists system_state (
                    key text primary key,
                    value text,
                    updated_at text default current_timestamp
                );
                create table if not exists item_groups (
                    id integer primary key,
                    name text not null unique,
                    description text default '',
                    is_active integer default 1,
                    created_at text default current_timestamp,
                    updated_at text default current_timestamp
                );
                create table if not exists item_catalog (
                    id integer primary key,
                    name text not null,
                    vat_rate text default '',
                    group_id integer,
                    notes text default '',
                    is_active integer default 1,
                    created_at text default current_timestamp,
                    updated_at text default current_timestamp,
                    foreign key(group_id) references item_groups(id)
                );
                create table if not exists document_pages (
                    id integer primary key,
                    file_id integer not null,
                    page_no integer not null,
                    width real default 0,
                    height real default 0,
                    rotation_deg real default 0,
                    text_layer_present integer default 0,
                    source_kind text default '',
                    ocr_status text default '',
                    confidence_avg real,
                    extracted_text text default '',
                    created_at text default current_timestamp,
                    unique(file_id, page_no),
                    foreign key(file_id) references files(id)
                );
                create table if not exists page_text_blocks (
                    id integer primary key,
                    page_id integer not null,
                    block_no integer not null,
                    raw_text text default '',
                    normalized_text text default '',
                    bbox_json text default '{}',
                    confidence real,
                    source_engine text default '',
                    source_kind text default '',
                    unique(page_id, block_no),
                    foreign key(page_id) references document_pages(id)
                );
                create table if not exists field_candidates (
                    id integer primary key,
                    document_id integer not null,
                    page_id integer,
                    field_name text not null,
                    raw_value text default '',
                    normalized_value text default '',
                    confidence real default 0,
                    source_kind text default '',
                    bbox_json text default '{}',
                    chosen integer default 0,
                    created_at text default current_timestamp,
                    foreign key(document_id) references processing_documents(id),
                    foreign key(page_id) references document_pages(id)
                );
                create table if not exists line_item_candidates (
                    id integer primary key,
                    document_id integer not null,
                    page_id integer,
                    line_no integer default 1,
                    name_raw text default '',
                    qty_raw text default '',
                    unit_price_raw text default '',
                    total_price_raw text default '',
                    vat_raw text default '',
                    confidence real default 0,
                    bbox_json text default '{}',
                    chosen integer default 0,
                    created_at text default current_timestamp,
                    foreign key(document_id) references processing_documents(id),
                    foreign key(page_id) references document_pages(id)
                );
                '''
            )
            self._ensure_column(conn, 'processing_documents', 'supplier_id', 'integer')
            self._ensure_column(conn, 'processing_documents', 'quarantine_reason', 'text default ""')
            self._ensure_column(conn, 'processing_documents', 'manual_payload_json', 'text default "{}"')
            self._ensure_column(conn, 'processing_documents', 'preview_page', 'integer default 1')
            self._ensure_column(conn, 'processing_suppliers', 'document_id', 'integer')
            self._ensure_column(conn, 'processing_suppliers', 'ares_payload_json', 'text default ""')
            self._ensure_column(conn, 'processing_suppliers', 'source_layer', 'text default "first_layer"')
            self._ensure_column(conn, 'visual_patterns', 'page_no', 'integer default 1')
            self._ensure_column(conn, 'visual_patterns', 'preview_state', 'text default "{}"')
            self._ensure_column(conn, 'visual_patterns', 'updated_at', 'text default current_timestamp')
            self._ensure_column(conn, 'processing_documents', 'promotion_status', 'text default ""')
            self._ensure_column(conn, 'processing_documents', 'promotion_blocked_reason', 'text default ""')
            self._ensure_column(conn, 'processing_documents', 'promoted_at', 'text')
            self._ensure_column(conn, 'processing_documents', 'segment_index', 'integer default 1')
            self._ensure_column(conn, 'processing_documents', 'page_from', 'integer default 1')
            self._ensure_column(conn, 'processing_documents', 'page_to', 'integer default 1')
            self._ensure_column(conn, 'processing_documents', 'split_status', 'text default "single"')
            self._ensure_column(conn, 'processing_documents', 'split_confidence', 'real default 1')
            self._ensure_column(conn, 'processing_items', 'source_kind', 'text default "automatic"')
            conn.execute("create table if not exists promotion_audit (id integer primary key, created_at text default current_timestamp, file_id integer, source_document_id integer, final_document_id integer, promotion_key text, outcome text, reason text, correlation_id text, payload_json text)")
            self._create_index_if_table_exists(conn, 'files', 'idx_files_sha256', 'create index if not exists idx_files_sha256 on files(sha256)')
            self._create_index_if_table_exists(conn, 'files', 'idx_files_status_id', 'create index if not exists idx_files_status_id on files(status, id)')
            self._create_index_if_table_exists(conn, 'processing_documents', 'idx_processing_documents_file_id', 'create index if not exists idx_processing_documents_file_id on processing_documents(file_id)')
            self._create_index_if_table_exists(conn, 'processing_documents', 'idx_processing_documents_status_state', 'create index if not exists idx_processing_documents_status_state on processing_documents(status, final_state, id)')
            self._create_index_if_table_exists(conn, 'processing_documents', 'idx_processing_documents_period', 'create index if not exists idx_processing_documents_period on processing_documents(substr(issued_at, 1, 7))')
            self._create_index_if_table_exists(conn, 'processing_suppliers', 'idx_processing_suppliers_document_id', 'create index if not exists idx_processing_suppliers_document_id on processing_suppliers(document_id)')
            self._create_index_if_table_exists(conn, 'processing_items', 'idx_processing_items_document_id', 'create index if not exists idx_processing_items_document_id on processing_items(document_id, id)')
            self._create_index_if_table_exists(conn, 'processing_items', 'idx_processing_items_vat_rate', 'create index if not exists idx_processing_items_vat_rate on processing_items(vat_rate, document_id)')
            self._create_index_if_table_exists(conn, 'document_pages', 'idx_document_pages_file_page', 'create index if not exists idx_document_pages_file_page on document_pages(file_id, page_no)')
            self._create_index_if_table_exists(conn, 'page_text_blocks', 'idx_page_text_blocks_page_block', 'create index if not exists idx_page_text_blocks_page_block on page_text_blocks(page_id, block_no)')
            self._create_index_if_table_exists(conn, 'field_candidates', 'idx_field_candidates_doc_field', 'create index if not exists idx_field_candidates_doc_field on field_candidates(document_id, field_name, chosen, confidence desc)')
            self._create_index_if_table_exists(conn, 'line_item_candidates', 'idx_line_item_candidates_doc', 'create index if not exists idx_line_item_candidates_doc on line_item_candidates(document_id, chosen, line_no)')
            self._create_index_if_table_exists(conn, 'extraction_attempts', 'idx_extraction_attempts_file_id', 'create index if not exists idx_extraction_attempts_file_id on extraction_attempts(file_id, id desc)')
            self._create_index_if_table_exists(conn, 'extraction_attempts', 'idx_extraction_attempts_document_id', 'create index if not exists idx_extraction_attempts_document_id on extraction_attempts(document_id, id desc)')
            self._create_index_if_table_exists(conn, 'document_results', 'idx_document_results_file_id', 'create index if not exists idx_document_results_file_id on document_results(file_id)')
            self._create_index_if_table_exists(conn, 'visual_patterns', 'idx_visual_patterns_active_name', 'create index if not exists idx_visual_patterns_active_name on visual_patterns(is_active, name)')
            self._create_index_if_table_exists(conn, 'item_groups', 'idx_item_groups_active_name', 'create index if not exists idx_item_groups_active_name on item_groups(is_active, name)')
            self._create_index_if_table_exists(conn, 'item_catalog', 'idx_item_catalog_group_active_name', 'create index if not exists idx_item_catalog_group_active_name on item_catalog(group_id, is_active, name)')
            self._create_index_if_table_exists(conn, 'ares_validations', 'idx_ares_validations_supplier_checked', 'create index if not exists idx_ares_validations_supplier_checked on ares_validations(supplier_ico, checked_at desc)')
            self._backfill_legacy_unrecognized_rows(conn)
            conn.commit()
        with self.prod_connection() as conn:
            conn.executescript(
                '''
                create virtual table if not exists final_documents_fts using fts5(document_number, vat_summary);
                create virtual table if not exists final_items_fts using fts5(name);
                create table if not exists reporting_cache_snapshots (
                    cache_key text primary key,
                    payload_json text not null,
                    updated_at text default current_timestamp
                );
                '''
            )
            self._ensure_column(conn, 'final_suppliers', 'ares_payload_json', 'text default ""')
            self._ensure_column(conn, 'final_suppliers', 'source_supplier_id', 'integer')
            self._ensure_column(conn, 'final_documents', 'source_file_id', 'integer')
            self._ensure_column(conn, 'final_documents', 'source_correlation_id', 'text')
            self._ensure_column(conn, 'final_documents', 'promotion_key', 'text')
            self._ensure_column(conn, 'final_documents', 'business_document_key', 'text')
            self._ensure_column(conn, 'final_items', 'source_item_id', 'integer')
            conn.execute('create unique index if not exists idx_final_suppliers_ico on final_suppliers(ico) where coalesce(ico, "") <> ""')
            conn.execute('create unique index if not exists idx_final_documents_source_document on final_documents(source_document_id)')
            conn.execute('create unique index if not exists idx_final_documents_promotion_key on final_documents(promotion_key) where coalesce(promotion_key, "") <> ""')
            conn.execute('create unique index if not exists idx_final_documents_business_key on final_documents(business_document_key) where coalesce(business_document_key, "") <> ""')
            conn.execute('create unique index if not exists idx_final_items_source_item on final_items(source_item_id) where source_item_id is not null')
            self._create_index_if_table_exists(conn, 'final_documents', 'idx_final_documents_issued_id', 'create index if not exists idx_final_documents_issued_id on final_documents(issued_at desc, id desc)')
            self._create_index_if_table_exists(conn, 'final_documents', 'idx_final_documents_supplier_issued', 'create index if not exists idx_final_documents_supplier_issued on final_documents(supplier_id, issued_at desc, id desc)')
            self._create_index_if_table_exists(conn, 'final_documents', 'idx_final_documents_period', 'create index if not exists idx_final_documents_period on final_documents(substr(issued_at, 1, 7), id desc)')
            self._create_index_if_table_exists(conn, 'final_items', 'idx_final_items_document_id', 'create index if not exists idx_final_items_document_id on final_items(document_id, id)')
            self._create_index_if_table_exists(conn, 'final_items', 'idx_final_items_vat_document', 'create index if not exists idx_final_items_vat_document on final_items(vat_rate, document_id)')
            self._create_index_if_table_exists(conn, 'final_suppliers', 'idx_final_suppliers_name', 'create index if not exists idx_final_suppliers_name on final_suppliers(name)')
            conn.commit()

    def _ensure_column(self, conn: sqlite3.Connection, table: str, column: str, definition: str) -> None:
        cols = {row[1] for row in conn.execute(f'pragma table_info({table})').fetchall()}
        if column not in cols:
            conn.execute(f'alter table {table} add column {column} {definition}')

    def _table_exists(self, conn: sqlite3.Connection, table: str) -> bool:
        row = conn.execute("select 1 from sqlite_master where type='table' and name=?", (table,)).fetchone()
        return row is not None

    def _create_index_if_table_exists(self, conn: sqlite3.Connection, table: str, _index_name: str, statement: str) -> None:
        if self._table_exists(conn, table):
            conn.execute(statement)

    def _review_state(self, state: str) -> str:
        return str(state or '').strip()

    def _legacy_unrecognized_reason_sql(self, reason_expr: str) -> str:
        return f'''
            lower(coalesce({reason_expr}, '')) like '%offline i openai cesty selhaly%'
            or lower(coalesce({reason_expr}, '')) like '%openai fallback nedodal validni vysledek%'
            or lower(coalesce({reason_expr}, '')) like '%openai vetev selhala%'
            or lower(coalesce({reason_expr}, '')) like '%doklad neobsahuje citelny text pro openai vetev%'
            or lower(coalesce({reason_expr}, '')) like '%chybi ulozeny openai api key%'
            or lower(coalesce({reason_expr}, '')) like '%nerozpozn%'
            or lower(coalesce({reason_expr}, '')) like '%nekompletni%'
            or lower(coalesce({reason_expr}, '')) like '%ocr_fallback_failed%'
        '''

    def _backfill_legacy_unrecognized_rows(self, conn: sqlite3.Connection) -> None:
        file_reason_expr = "coalesce(nullif(processing_documents.quarantine_reason, ''), nullif(files.last_error, ''), '')"
        doc_reason_expr = "coalesce(nullif(quarantine_reason, ''), '')"
        result_reason_expr = "coalesce(reason, '')"
        conn.execute(
            f'''
            update files
            set status='unrecognized'
            where lower(coalesce(status, '')) = 'quarantine'
              and exists (
                  select 1
                  from processing_documents
                  where processing_documents.file_id = files.id
                    and ({self._legacy_unrecognized_reason_sql(file_reason_expr)})
              )
            '''
        )
        conn.execute(
            f'''
            update processing_documents
            set status='unrecognized',
                final_state='unrecognized'
            where lower(coalesce(status, final_state, '')) = 'quarantine'
              and ({self._legacy_unrecognized_reason_sql(doc_reason_expr)})
            '''
        )
        conn.execute(
            f'''
            update document_results
            set final_state='unrecognized'
            where lower(coalesce(final_state, '')) = 'quarantine'
              and ({self._legacy_unrecognized_reason_sql(result_reason_expr)})
            '''
        )

    def _fts_query(self, search: str) -> str:
        tokens = [token for token in re.findall(r'\w+', search.lower()) if token]
        return ' '.join(f'{token}*' for token in tokens)

    def _qualified(self, schema: str, table: str) -> str:
        return table if schema in {'', 'main'} else f'{schema}.{table}'

    def _sync_final_document_fts(self, conn: sqlite3.Connection, document_id: int, *, schema: str = 'main') -> None:
        final_documents_fts = self._qualified(schema, 'final_documents_fts')
        final_documents = self._qualified(schema, 'final_documents')
        conn.execute(f'delete from {final_documents_fts} where rowid=?', (document_id,))
        row = conn.execute(f'select document_number, vat_summary from {final_documents} where id=?', (document_id,)).fetchone()
        if row:
            conn.execute(
                f'insert into {final_documents_fts}(rowid, document_number, vat_summary) values (?, ?, ?)',
                (document_id, row['document_number'], row['vat_summary']),
            )

    def _sync_final_item_fts(self, conn: sqlite3.Connection, item_id: int, *, schema: str = 'main') -> None:
        final_items_fts = self._qualified(schema, 'final_items_fts')
        final_items = self._qualified(schema, 'final_items')
        conn.execute(f'delete from {final_items_fts} where rowid=?', (item_id,))
        row = conn.execute(f'select name from {final_items} where id=?', (item_id,)).fetchone()
        if row:
            conn.execute(f'insert into {final_items_fts}(rowid, name) values (?, ?)', (item_id, row['name']))

    def _sync_working_document_fts(self, conn: sqlite3.Connection, file_id: int) -> None:
        conn.execute('delete from working_documents_fts where rowid=?', (file_id,))
        row = conn.execute(
            '''
            select f.original_name, pd.document_number, ps.ico as supplier_ico, ps.name as supplier_name,
                   pd.quarantine_reason, f.last_error, f.correlation_id
            from files f
            join processing_documents pd on pd.file_id = f.id
            left join processing_suppliers ps on ps.document_id = pd.id
            where f.id=?
            ''',
            (file_id,),
        ).fetchone()
        if row:
            conn.execute(
                '''
                insert into working_documents_fts(
                    rowid, original_name, document_number, supplier_ico, supplier_name,
                    quarantine_reason, last_error, correlation_id
                ) values (?, ?, ?, ?, ?, ?, ?, ?)
                ''',
                (
                    file_id,
                    row['original_name'],
                    row['document_number'],
                    row['supplier_ico'],
                    row['supplier_name'],
                    row['quarantine_reason'],
                    row['last_error'],
                    row['correlation_id'],
                ),
            )

    def _sync_working_document_fts_by_document(self, conn: sqlite3.Connection, document_id: int) -> None:
        row = conn.execute('select file_id from processing_documents where id=?', (document_id,)).fetchone()
        if row:
            self._sync_working_document_fts(conn, int(row['file_id']))

    def _sync_attempt_fts(self, conn: sqlite3.Connection, attempt_id: int) -> None:
        conn.execute('delete from extraction_attempts_fts where rowid=?', (attempt_id,))
        row = conn.execute(
            '''
            select ea.id, f.original_name, ea.reason, ea.branch, ea.processor_name, ea.correlation_id
            from extraction_attempts ea
            join files f on f.id = ea.file_id
            where ea.id=?
            ''',
            (attempt_id,),
        ).fetchone()
        if row:
            conn.execute(
                '''
                insert into extraction_attempts_fts(
                    rowid, original_name, reason, branch, processor_name, correlation_id
                ) values (?, ?, ?, ?, ?, ?)
                ''',
                (
                    attempt_id,
                    row['original_name'],
                    row['reason'],
                    row['branch'],
                    row['processor_name'],
                    row['correlation_id'],
                ),
            )

    def _sha256(self, source_path: Path) -> str:
        digest = hashlib.sha256()
        with source_path.open('rb') as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b''):
                digest.update(chunk)
        return digest.hexdigest()

    def _safe_file_name(self, value: str) -> str:
        return ''.join(ch if ch.isalnum() or ch in '._-' else '_' for ch in value)

    def metrics(self) -> dict[str, int]:
        with self.work_connection() as work, self.prod_connection() as prod:
            queue_size = work.execute("select count(*) from files where status in ('new','queued','processing','retry_pending')").fetchone()[0]
            quarantine = work.execute("select count(*) from files where status='quarantine'").fetchone()[0]
            unrecognized = work.execute("select count(*) from files where status='unrecognized'").fetchone()[0]
            duplicates = work.execute("select count(*) from files where status = 'duplicate'").fetchone()[0]
            finals = prod.execute('select count(*) from final_documents').fetchone()[0]
            attempts = work.execute('select count(*) from extraction_attempts').fetchone()[0]
            suppliers = prod.execute('select count(*) from final_suppliers').fetchone()[0]
            return {'queue_size': int(queue_size), 'quarantine': int(quarantine), 'unrecognized': int(unrecognized), 'duplicates': int(duplicates), 'final_documents': int(finals), 'attempts': int(attempts), 'suppliers': int(suppliers)}

    def import_file_to_project(self, source_path: Path, *, reject_duplicates: bool = True, loaded_document: LoadedDocument | None = None, output_directory: Path | None = None) -> int:
        correlation_id = uuid.uuid4().hex
        sha256 = self._sha256(source_path)
        page_count = max(1, (loaded_document.page_count if loaded_document is not None else self.document_loader.inspect_page_count(source_path)))
        if reject_duplicates:
            with self.work_connection() as conn:
                duplicate = conn.execute('select id, status from files where sha256=? order by id desc limit 1', (sha256,)).fetchone()
                if duplicate:
                    if output_directory:
                        self._copy_external_output(source_path, output_directory, bucket='quarantine')
                    self._remove_consumed_input(source_path)
                    self.log_runtime(f'IMPORT duplicate-skip {source_path.name} -> file_id={duplicate["id"]} sha256={sha256[:10]}')
                    self.log_event('warning', 'import.duplicate', 'Soubor se stejným obsahem už v projektu existuje.', correlation_id, file_id=int(duplicate['id']), payload={'sha256_prefix': sha256[:12], 'source_name': source_path.name})
                    return int(duplicate['id'])

        stored_name = f'{datetime.now(UTC).strftime("%Y%m%d%H%M%S")}_{self._safe_file_name(source_path.name)}'
        target = self.project_path / 'documents' / 'originals' / stored_name
        target.parent.mkdir(parents=True, exist_ok=True)
        temp_target = target.with_suffix(f'{target.suffix}.partial')
        shutil.copy2(source_path, temp_target)
        temp_target.replace(target)
        try:
            with self.work_connection() as conn:
                cur = conn.execute(
                    '''insert into files (original_name, source_import, internal_path, sha256, file_type, size_bytes, page_count, received_at, source_type, status, correlation_id, last_error, imported_at)
                       values (?, ?, ?, ?, ?, ?, ?, ?, 'filesystem', 'queued', ?, '', ?)''',
                    (source_path.name, str(source_path), str(target.relative_to(self.project_path)), sha256, source_path.suffix.lower().lstrip('.'), source_path.stat().st_size, page_count, utc_now(), correlation_id, utc_now()),
                )
                file_id = int(cur.lastrowid)
                dcur = conn.execute("insert into processing_documents (file_id, document_type, family, status, final_state) values (?, 'expense', 'expense', 'draft', 'pending')", (file_id,))
                document_id = int(dcur.lastrowid)
                conn.execute('insert into processing_suppliers (document_id, ico, dic, name, vat_payer, address, ares_status, ares_payload_json, source_layer) values (?, "", "", "", 0, "", "pending", "{}", "first_layer")', (document_id,))
                self._sync_working_document_fts(conn, file_id)
                self._set_queue_size(conn)
                conn.commit()
        except Exception:
            try:
                if target.exists():
                    target.unlink()
            except Exception:
                pass
            raise
        self._remove_consumed_input(source_path)
        self.log_runtime(f'IMPORT {source_path.name} -> file_id={file_id} correlation_id={correlation_id}')
        self.log_event('info', 'import.created', 'Soubor byl zařazen do fronty.', correlation_id, file_id=file_id, document_id=document_id, payload={'sha256_prefix': sha256[:12]})
        return file_id

    def import_file(self, source_path: Path, *, reject_duplicates: bool = True, loaded_document: LoadedDocument | None = None, output_directory: Path | None = None) -> int:
        return self.import_file_to_project(source_path, reject_duplicates=reject_duplicates, loaded_document=loaded_document, output_directory=output_directory)

    def count_file_attempts(self, file_id: int, *, attempt_types: list[str] | None = None) -> int:
        sql = 'select count(*) from extraction_attempts where file_id=?'
        params: list[Any] = [file_id]
        if attempt_types:
            placeholders = ', '.join('?' for _ in attempt_types)
            sql += f' and attempt_type in ({placeholders})'
            params.extend(attempt_types)
        with self.work_connection() as conn:
            return int(conn.execute(sql, params).fetchone()[0])

    def count_document_attempts(self, document_id: int, *, attempt_types: list[str] | None = None) -> int:
        sql = 'select count(*) from extraction_attempts where document_id=?'
        params: list[Any] = [document_id]
        if attempt_types:
            placeholders = ', '.join('?' for _ in attempt_types)
            sql += f' and attempt_type in ({placeholders})'
            params.extend(attempt_types)
        with self.work_connection() as conn:
            return int(conn.execute(sql, params).fetchone()[0])

    def processing_document_page_for_file(self, file_id: int, *, page_size: int = 100) -> int:
        page_size = max(1, min(int(page_size or 100), 500))
        with self.work_connection() as conn:
            position = int(conn.execute('select count(*) from processing_documents where file_id>?', (file_id,)).fetchone()[0])
        return (position // page_size) + 1

    def final_document_page_for_id(self, document_id: int, *, page_size: int = 100) -> int:
        page_size = max(1, min(int(page_size or 100), 500))
        with self.prod_connection() as conn:
            row = conn.execute('select issued_at from final_documents where id=?', (document_id,)).fetchone()
            if not row:
                return 1
            position = int(
                conn.execute(
                    '''
                    select count(*)
                    from final_documents
                    where issued_at > ?
                       or (issued_at = ? and id > ?)
                    ''',
                    (row['issued_at'], row['issued_at'], document_id),
                ).fetchone()[0]
            )
        return (position // page_size) + 1

    def final_item_page_for_id(self, item_id: int, *, page_size: int = 200) -> int:
        page_size = max(1, min(int(page_size or 200), 500))
        with self.prod_connection() as conn:
            row = conn.execute(
                '''
                select fi.id, fd.issued_at
                from final_items fi
                join final_documents fd on fd.id=fi.document_id
                where fi.id=?
                ''',
                (item_id,),
            ).fetchone()
            if not row:
                return 1
            position = int(
                conn.execute(
                    '''
                    select count(*)
                    from final_items fi
                    join final_documents fd on fd.id=fi.document_id
                    where fd.issued_at > ?
                       or (fd.issued_at = ? and fi.id > ?)
                    ''',
                    (row['issued_at'], row['issued_at'], item_id),
                ).fetchone()[0]
            )
        return (position // page_size) + 1

    def get_pending_files(self) -> list[sqlite3.Row]:
        with self.work_connection() as conn:
            return conn.execute('''select f.*, pd.id as document_id from files f join processing_documents pd on pd.file_id=f.id where f.status in ('queued','retry_pending') order by f.id asc''').fetchall()

    def get_file_record(self, file_id: int) -> sqlite3.Row | None:
        with self.work_connection() as conn:
            return conn.execute('''select f.*, pd.id as document_id from files f join processing_documents pd on pd.file_id=f.id where f.id = ? order by coalesce(pd.segment_index, 1) asc, pd.id asc''', (file_id,)).fetchone()

    def list_processing_documents_for_file(self, file_id: int) -> list[sqlite3.Row]:
        with self.work_connection() as conn:
            return conn.execute(
                '''
                select pd.*, ps.ico, ps.name as supplier_name, ps.dic, ps.vat_payer, ps.address, ps.ares_status, ps.ares_payload_json
                from processing_documents pd
                left join processing_suppliers ps on ps.document_id = pd.id
                where pd.file_id=?
                order by coalesce(pd.segment_index, 1) asc, pd.id asc
                ''',
                (file_id,),
            ).fetchall()

    def sync_processing_document_segments(self, file_id: int, segments: list[dict[str, Any]]) -> list[int]:
        normalized_segments = list(segments or [])
        if not normalized_segments:
            normalized_segments = [{'segment_index': 1, 'page_from': 1, 'page_to': 1, 'split_status': 'single', 'split_confidence': 1.0}]
        with self.work_connection() as conn:
            existing = conn.execute('select id from processing_documents where file_id=? order by coalesce(segment_index, 1) asc, id asc', (file_id,)).fetchall()
            if conn.execute('select count(*) from extraction_attempts where file_id=?', (file_id,)).fetchone()[0]:
                return [int(row['id']) for row in existing]
            for row in existing:
                document_id = int(row['id'])
                conn.execute('delete from processing_items where document_id=?', (document_id,))
                conn.execute('delete from processing_suppliers where document_id=?', (document_id,))
                conn.execute('delete from field_candidates where document_id=?', (document_id,))
                conn.execute('delete from line_item_candidates where document_id=?', (document_id,))
            conn.execute('delete from processing_documents where file_id=?', (file_id,))
            document_ids: list[int] = []
            for index, segment in enumerate(normalized_segments, start=1):
                cur = conn.execute(
                    '''
                    insert into processing_documents (
                        file_id, document_type, family, status, final_state, segment_index, page_from, page_to, split_status, split_confidence
                    ) values (?, 'expense', 'expense', 'draft', 'pending', ?, ?, ?, ?, ?)
                    ''',
                    (
                        file_id,
                        int(segment.get('segment_index') or index),
                        int(segment.get('page_from') or 1),
                        int(segment.get('page_to') or segment.get('page_from') or 1),
                        str(segment.get('split_status') or ('single' if len(normalized_segments) == 1 else 'split')),
                        float(segment.get('split_confidence') or 1.0),
                    ),
                )
                document_id = int(cur.lastrowid or 0)
                document_ids.append(document_id)
                conn.execute(
                    '''
                    insert into processing_suppliers (
                        document_id, ico, dic, name, vat_payer, address, ares_status, ares_payload_json, source_layer
                    ) values (?, "", "", "", 0, "", "pending", "{}", "first_layer")
                    ''',
                    (document_id,),
                )
            self._sync_working_document_fts(conn, file_id)
            conn.commit()
        return document_ids

    def update_file_page_count(self, file_id: int, page_count: int) -> None:
        with self.work_connection() as conn:
            conn.execute('update files set page_count=? where id=?', (int(page_count), file_id))
            conn.commit()

    def replace_document_pages(self, file_id: int, pages: list[dict[str, Any]]) -> list[int]:
        with self.work_connection() as conn:
            page_rows = conn.execute('select id from document_pages where file_id=?', (file_id,)).fetchall()
            page_ids_to_replace = [int(row['id']) for row in page_rows]
            for page_id in page_ids_to_replace:
                conn.execute('delete from field_candidates where page_id=?', (page_id,))
                conn.execute('delete from line_item_candidates where page_id=?', (page_id,))
            for row in page_rows:
                conn.execute('delete from page_text_blocks where page_id=?', (int(row['id']),))
            conn.execute('delete from document_pages where file_id=?', (file_id,))
            page_ids: list[int] = []
            for page in pages:
                cur = conn.execute(
                    '''
                    insert into document_pages (
                        file_id, page_no, width, height, rotation_deg, text_layer_present,
                        source_kind, ocr_status, confidence_avg, extracted_text, created_at
                    ) values (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ''',
                    (
                        file_id,
                        int(page.get('page_no') or 1),
                        float(page.get('width') or 0),
                        float(page.get('height') or 0),
                        float(page.get('rotation_deg') or 0),
                        int(bool(page.get('text_layer_present'))),
                        str(page.get('source_kind') or ''),
                        str(page.get('ocr_status') or ''),
                        page.get('confidence_avg'),
                        str(page.get('extracted_text') or ''),
                        utc_now(),
                    ),
                )
                page_ids.append(int(cur.lastrowid or 0))
            conn.commit()
            return page_ids

    def replace_page_text_blocks(self, page_id: int, blocks: list[dict[str, Any]]) -> None:
        with self.work_connection() as conn:
            conn.execute('delete from page_text_blocks where page_id=?', (page_id,))
            for block in blocks:
                conn.execute(
                    '''
                    insert into page_text_blocks (
                        page_id, block_no, raw_text, normalized_text, bbox_json, confidence, source_engine, source_kind
                    ) values (?, ?, ?, ?, ?, ?, ?, ?)
                    ''',
                    (
                        page_id,
                        int(block.get('block_no') or 1),
                        str(block.get('raw_text') or ''),
                        str(block.get('normalized_text') or ''),
                        json.dumps(block.get('bbox') or {}, ensure_ascii=False),
                        block.get('confidence'),
                        str(block.get('source_engine') or ''),
                        str(block.get('source_kind') or ''),
                    ),
                )
            conn.commit()

    def list_document_pages(self, file_id: int) -> list[sqlite3.Row]:
        with self.work_connection() as conn:
            return conn.execute('select * from document_pages where file_id=? order by page_no asc', (file_id,)).fetchall()

    def list_page_text_blocks(self, file_id: int) -> list[sqlite3.Row]:
        with self.work_connection() as conn:
            return conn.execute(
                '''
                select b.*, p.file_id, p.page_no
                from page_text_blocks b
                join document_pages p on p.id = b.page_id
                where p.file_id=?
                order by p.page_no asc, b.block_no asc
                ''',
                (file_id,),
            ).fetchall()

    def replace_field_candidates(self, document_id: int, candidates: list[dict[str, Any]]) -> None:
        with self.work_connection() as conn:
            conn.execute('delete from field_candidates where document_id=?', (document_id,))
            for candidate in candidates:
                conn.execute(
                    '''
                    insert into field_candidates (
                        document_id, page_id, field_name, raw_value, normalized_value, confidence, source_kind, bbox_json, chosen, created_at
                    ) values (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ''',
                    (
                        document_id,
                        candidate.get('page_id'),
                        str(candidate.get('field_name') or ''),
                        str(candidate.get('raw_value') or ''),
                        str(candidate.get('normalized_value') or ''),
                        float(candidate.get('confidence') or 0),
                        str(candidate.get('source_kind') or ''),
                        json.dumps(candidate.get('bbox') or {}, ensure_ascii=False),
                        int(bool(candidate.get('chosen'))),
                        utc_now(),
                    ),
                )
            conn.commit()

    def list_field_candidates(self, document_id: int, field_name: str = '') -> list[sqlite3.Row]:
        sql = 'select * from field_candidates where document_id=?'
        params: list[Any] = [document_id]
        if field_name:
            sql += ' and field_name=?'
            params.append(field_name)
        sql += ' order by field_name asc, chosen desc, confidence desc, id asc'
        with self.work_connection() as conn:
            return conn.execute(sql, params).fetchall()

    def replace_line_item_candidates(self, document_id: int, candidates: list[dict[str, Any]]) -> None:
        with self.work_connection() as conn:
            conn.execute('delete from line_item_candidates where document_id=?', (document_id,))
            for candidate in candidates:
                conn.execute(
                    '''
                    insert into line_item_candidates (
                        document_id, page_id, line_no, name_raw, qty_raw, unit_price_raw, total_price_raw, vat_raw, confidence, bbox_json, chosen, created_at
                    ) values (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ''',
                    (
                        document_id,
                        candidate.get('page_id'),
                        int(candidate.get('line_no') or 1),
                        str(candidate.get('name_raw') or ''),
                        str(candidate.get('qty_raw') or ''),
                        str(candidate.get('unit_price_raw') or ''),
                        str(candidate.get('total_price_raw') or ''),
                        str(candidate.get('vat_raw') or ''),
                        float(candidate.get('confidence') or 0),
                        json.dumps(candidate.get('bbox') or {}, ensure_ascii=False),
                        int(bool(candidate.get('chosen'))),
                        utc_now(),
                    ),
                )
            conn.commit()

    def list_line_item_candidates(self, document_id: int) -> list[sqlite3.Row]:
        with self.work_connection() as conn:
            return conn.execute('select * from line_item_candidates where document_id=? order by chosen desc, line_no asc, id asc', (document_id,)).fetchall()

    def mark_file_status(self, file_id: int, status: str, last_error: str = '') -> None:
        status = self._review_state(status)
        with self.work_connection() as conn:
            conn.execute('update files set status=?, last_error=? where id=?', (status, last_error, file_id))
            if status in {'processing', 'retry_pending', 'quarantine', 'unrecognized'}:
                conn.execute(
                    '''
                    update processing_documents
                    set status=?,
                        final_state=?,
                        quarantine_reason=case
                            when ? in ('quarantine', 'unrecognized') and ? <> '' then ?
                            else quarantine_reason
                        end
                    where file_id=?
                    ''',
                    (status, status, status, last_error, last_error, file_id),
                )
            self._sync_working_document_fts(conn, file_id)
            self._set_queue_size(conn)
            conn.commit()

    def update_processing_document(self, document_id: int, *, document_number: str | None = None, issued_at: str | None = None, total_with_vat: float | None = None, quarantine_reason: str | None = None, manual_payload: dict[str, Any] | None = None) -> None:
        parts = []
        params: list[Any] = []
        if document_number is not None:
            parts.append('document_number=?'); params.append(document_number)
        if issued_at is not None:
            parts.append('issued_at=?'); params.append(issued_at)
        if total_with_vat is not None:
            parts.append('total_with_vat=?'); params.append(total_with_vat)
        if quarantine_reason is not None:
            parts.append('quarantine_reason=?'); params.append(quarantine_reason)
        if manual_payload is not None:
            parts.append('manual_payload_json=?'); params.append(json.dumps(manual_payload, ensure_ascii=False))
        if not parts:
            return
        params.append(document_id)
        with self.work_connection() as conn:
            conn.execute(f'update processing_documents set {", ".join(parts)} where id=?', params)
            self._sync_working_document_fts_by_document(conn, document_id)
            conn.commit()

    def update_processing_supplier(self, document_id: int, *, ico: str, name: str = '', dic: str = '', vat_payer: bool = False, address: str = '', ares_status: str = 'pending', ares_payload: dict[str, Any] | None = None) -> None:
        with self.work_connection() as conn:
            row = conn.execute('select id from processing_suppliers where document_id=?', (document_id,)).fetchone()
            payload_json = json.dumps(ares_payload or {}, ensure_ascii=False)
            if row:
                conn.execute('''update processing_suppliers set ico=?, name=?, dic=?, vat_payer=?, address=?, ares_status=?, ares_payload_json=? where document_id=?''', (ico, name, dic, int(vat_payer), address, ares_status, payload_json, document_id))
                supplier_id = int(row['id'])
            else:
                cur = conn.execute('''insert into processing_suppliers (document_id, ico, dic, name, vat_payer, address, ares_status, ares_payload_json, source_layer) values (?, ?, ?, ?, ?, ?, ?, ?, 'first_layer')''', (document_id, ico, dic, name, int(vat_payer), address, ares_status, payload_json))
                supplier_id = int(cur.lastrowid)
            conn.execute('update processing_documents set supplier_id=? where id=?', (supplier_id, document_id))
            self._sync_working_document_fts_by_document(conn, document_id)
            conn.commit()

    def create_attempt(self, file_id: int, document_id: int, attempt_type: str, branch: str, processor_name: str, correlation_id: str) -> int:
        with self.work_connection() as conn:
            order_no = conn.execute('select coalesce(max(attempt_order),0)+1 from extraction_attempts where file_id=?', (file_id,)).fetchone()[0]
            cur = conn.execute(
                '''insert into extraction_attempts (file_id, document_id, attempt_order, attempt_type, started_at, branch, result, reason, blocking_error, next_step, document_state, correlation_id, processor_name, payload_json)
                   values (?, ?, ?, ?, ?, ?, 'running', '', 0, '', 'processing', ?, ?, '{}')''',
                (file_id, document_id, order_no, attempt_type, utc_now(), branch, correlation_id, processor_name),
            )
            attempt_id = int(cur.lastrowid)
            conn.execute('update processing_documents set last_attempt_id=? where id=?', (attempt_id, document_id))
            self._sync_attempt_fts(conn, attempt_id)
            conn.commit()
        self.log_event('info', 'attempt.started', f'Start pokusu {attempt_type}.', correlation_id, file_id=file_id, document_id=document_id, attempt_id=attempt_id)
        return attempt_id

    def finish_attempt(self, attempt_id: int, *, result: str, reason: str, next_step: str, document_state: str, blocking_error: bool = False, payload: dict[str, Any] | None = None) -> None:
        with self.work_connection() as conn:
            row = conn.execute('select * from extraction_attempts where id=?', (attempt_id,)).fetchone()
            conn.execute('''update extraction_attempts set finished_at=?, result=?, reason=?, next_step=?, document_state=?, blocking_error=?, payload_json=? where id=?''', (utc_now(), result, reason, next_step, document_state, int(blocking_error), json.dumps(payload or {}, ensure_ascii=False), attempt_id))
            conn.execute('update processing_documents set status=?, final_state=? where id=?', (document_state, document_state, row['document_id']))
            self._sync_attempt_fts(conn, attempt_id)
            conn.commit()
        self.log_event('info' if result == 'success' else 'warning', 'attempt.finished', reason, row['correlation_id'], file_id=row['file_id'], document_id=row['document_id'], attempt_id=attempt_id, payload=payload)


    def _processing_snapshot(self, conn: sqlite3.Connection, document_id: int) -> tuple[sqlite3.Row | None, sqlite3.Row | None]:
        doc = conn.execute('select * from processing_documents where id=?', (document_id,)).fetchone()
        supplier = conn.execute('select * from processing_suppliers where document_id=?', (document_id,)).fetchone()
        return doc, supplier

    def replace_processing_items(self, document_id: int, items: list[dict[str, Any]]) -> None:
        with self.work_connection() as conn:
            conn.execute('delete from processing_items where document_id=?', (document_id,))
            for item in items:
                conn.execute(
                    "insert into processing_items (document_id, name, quantity, unit_price, total_price, vat_rate, source_kind) values (?, ?, ?, ?, ?, ?, ?)",
                    (
                        document_id,
                        str(item.get('name', '')).strip(),
                        float(item.get('quantity') or 1),
                        float(item.get('unit_price') or item.get('total_price') or 0),
                        float(item.get('total_price') or 0),
                        str(item.get('vat_rate') or '').strip(),
                        str(item.get('source_kind') or 'automatic'),
                    ),
                )
            conn.commit()

    def list_processing_items(self, document_id: int) -> list[sqlite3.Row]:
        with self.work_connection() as conn:
            return conn.execute('select * from processing_items where document_id=? order by id asc', (document_id,)).fetchall()

    def _promotion_key(self, pd: sqlite3.Row, ps: sqlite3.Row) -> str:
        return f"{ps['ico'] or 'noico'}::{pd['document_number'] or 'nodoc'}::{pd['issued_at'] or 'nodate'}::{float(pd['total_with_vat'] or 0):.2f}"

    def _business_document_key(self, pd: sqlite3.Row, ps: sqlite3.Row) -> str:
        return self._promotion_key(pd, ps)

    def _upsert_final_supplier(self, conn: sqlite3.Connection, ps: sqlite3.Row, *, schema: str = 'main') -> int:
        final_suppliers = self._qualified(schema, 'final_suppliers')
        supplier_row = conn.execute(f'select id from {final_suppliers} where ico=?', (ps['ico'],)).fetchone()
        if supplier_row:
            supplier_id = int(supplier_row['id'])
            conn.execute(
                f'''update {final_suppliers}
                    set dic=?, name=?, vat_payer=?, address=?, ares_checked_at=?, ares_payload_json=?, source_supplier_id=?
                    where id=?''',
                (ps['dic'], ps['name'], ps['vat_payer'], ps['address'], utc_now(), ps['ares_payload_json'], ps['id'], supplier_id),
            )
            return supplier_id
        cur = conn.execute(
            f'''insert into {final_suppliers}
                (ico, dic, name, vat_payer, address, ares_checked_at, ares_payload_json, source_supplier_id)
                values (?, ?, ?, ?, ?, ?, ?, ?)''',
            (ps['ico'], ps['dic'], ps['name'] or 'Neznámý dodavatel', ps['vat_payer'], ps['address'], utc_now(), ps['ares_payload_json'], ps['id']),
        )
        return int(cur.lastrowid or 0)

    def _final_snapshot_matches(
        self,
        conn: sqlite3.Connection,
        final_document_id: int,
        *,
        supplier_id: int,
        pd: sqlite3.Row,
        items: list[sqlite3.Row],
        vat_summary: str,
        item_total: float,
        total_without_vat: float,
        original_path: str,
        schema: str = 'main',
    ) -> bool:
        final_documents = self._qualified(schema, 'final_documents')
        final_items = self._qualified(schema, 'final_items')
        current_doc = conn.execute(
            f'''select supplier_id, issued_at, document_number, total_without_vat, total_with_vat, vat_summary, original_file_path
                from {final_documents}
                where id=?''',
            (final_document_id,),
        ).fetchone()
        if not current_doc:
            return False
        expected_doc = {
            'supplier_id': int(supplier_id),
            'issued_at': str(pd['issued_at'] or ''),
            'document_number': str(pd['document_number'] or ''),
            'total_without_vat': round(float(total_without_vat or 0), 2),
            'total_with_vat': round(float(item_total or 0), 2),
            'vat_summary': str(vat_summary or ''),
            'original_file_path': str(original_path or ''),
        }
        current_doc_data = {
            'supplier_id': int(current_doc['supplier_id'] or 0),
            'issued_at': str(current_doc['issued_at'] or ''),
            'document_number': str(current_doc['document_number'] or ''),
            'total_without_vat': round(float(current_doc['total_without_vat'] or 0), 2),
            'total_with_vat': round(float(current_doc['total_with_vat'] or 0), 2),
            'vat_summary': str(current_doc['vat_summary'] or ''),
            'original_file_path': str(current_doc['original_file_path'] or ''),
        }
        if current_doc_data != expected_doc:
            return False
        current_items = [
            {
                'source_item_id': int(row['source_item_id']),
                'name': str(row['name'] or ''),
                'quantity': round(float(row['quantity'] or 0), 4),
                'unit_price': round(float(row['unit_price'] or 0), 4),
                'total_price': round(float(row['total_price'] or 0), 4),
                'vat_rate': str(row['vat_rate'] or ''),
            }
            for row in conn.execute(
                f'''select source_item_id, name, quantity, unit_price, total_price, vat_rate
                    from {final_items}
                    where document_id=?
                    order by id asc''',
                (final_document_id,),
            ).fetchall()
        ]
        expected_items = [
            {
                'source_item_id': int(item['id']),
                'name': str(item['name'] or ''),
                'quantity': round(float(item['quantity'] or 0), 4),
                'unit_price': round(float(item['unit_price'] or 0), 4),
                'total_price': round(float(item['total_price'] or 0), 4),
                'vat_rate': str(item['vat_rate'] or ''),
            }
            for item in items
        ]
        return current_items == expected_items

    def _replace_final_items(self, conn: sqlite3.Connection, final_document_id: int, items: list[sqlite3.Row], *, schema: str = 'main') -> None:
        final_items = self._qualified(schema, 'final_items')
        final_items_fts = self._qualified(schema, 'final_items_fts')
        existing_ids = conn.execute(f'select id from {final_items} where document_id=?', (final_document_id,)).fetchall()
        for row in existing_ids:
            conn.execute(f'delete from {final_items_fts} where rowid=?', (int(row['id']),))
        conn.execute(f'delete from {final_items} where document_id=?', (final_document_id,))
        for item in items:
            icur = conn.execute(
                f'''insert into {final_items} (document_id, source_item_id, name, quantity, unit_price, total_price, vat_rate)
                    values (?, ?, ?, ?, ?, ?, ?)''',
                (final_document_id, item['id'], item['name'], item['quantity'], item['unit_price'], item['total_price'], item['vat_rate']),
            )
            self._sync_final_item_fts(conn, int(icur.lastrowid or 0), schema=schema)

    def _normalize_ico(self, ico: str) -> str:
        return ''.join(ch for ch in str(ico or '').strip() if ch.isdigit())

    def _normalize_issued_at(self, issued_at: str) -> str:
        value = str(issued_at or '').strip()
        if re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
            return value
        match = re.fullmatch(r'(\d{1,2})[./-](\d{1,2})[./-](\d{4})', value)
        if not match:
            raise ValueError('Datum dokladu musi byt ve formatu YYYY-MM-DD nebo DD.MM.RRRR.')
        day, month, year = match.groups()
        return f'{int(year):04d}-{int(month):02d}-{int(day):02d}'

    def _parse_vat_rate(self, value: Any) -> float | None:
        normalized = str(value or '').strip().replace('%', '').replace(',', '.')
        if not normalized:
            return None
        try:
            parsed = float(normalized)
        except ValueError:
            return None
        return parsed if parsed >= 0 else None

    def _validate_processing_items(self, items: list[Any], *, expected_total: float | None = None) -> str:
        if not items:
            return 'Doklad musi obsahovat alespon jednu polozku.'
        computed_total = 0.0
        for item in items:
            name = str(item['name'] if isinstance(item, sqlite3.Row) else item.get('name', '')).strip()
            vat_rate = str(item['vat_rate'] if isinstance(item, sqlite3.Row) else item.get('vat_rate', '')).strip()
            try:
                total_price = round(float(item['total_price'] if isinstance(item, sqlite3.Row) else item.get('total_price') or 0), 2)
            except (TypeError, ValueError):
                return 'Polozky dokladu musi obsahovat validni castky.'
            if not name:
                return 'Kazda polozka musi mit nazev.'
            if total_price <= 0:
                return 'Kazda polozka musi mit kladnou cenu celkem.'
            if not vat_rate:
                return 'Kazda polozka musi mit druh DPH.'
            if self._parse_vat_rate(vat_rate) is None:
                return 'Druh DPH na polozkach musi byt validni sazba.'
            computed_total = round(computed_total + total_price, 2)
        if expected_total is not None and round(float(expected_total or 0), 2) != computed_total:
            return 'Soucet polozek neodpovida celkove castce dokladu.'
        return ''

    def _vat_summary_from_items(self, items: list[sqlite3.Row]) -> tuple[str, float, float]:
        buckets: dict[str, float] = {}
        gross_total = 0.0
        net_total = 0.0
        for item in items:
            amount = round(float(item['total_price'] or 0), 2)
            gross_total = round(gross_total + amount, 2)
            rate = str(item['vat_rate'] or '').strip() or '0'
            parsed_rate = self._parse_vat_rate(rate)
            if parsed_rate is None:
                raise ValueError('Druh DPH na polozkach musi byt validni sazba.')
            divisor = 1 + (parsed_rate / 100)
            net_total = round(net_total + (amount if divisor <= 0 else amount / divisor), 2)
            buckets[rate] = round(buckets.get(rate, 0.0) + amount, 2)
        summary = ', '.join(f"{rate}%: {amount:.2f}" for rate, amount in sorted(buckets.items())) or '0%: 0.00'
        return summary, round(gross_total, 2), round(net_total, 2)

    def record_promotion_audit(self, file_id: int, document_id: int, final_document_id: int | None, promotion_key: str, outcome: str, reason: str, correlation_id: str, payload: dict[str, Any] | None = None) -> None:
        with self.work_connection() as conn:
            conn.execute("insert into promotion_audit (file_id, source_document_id, final_document_id, promotion_key, outcome, reason, correlation_id, payload_json) values (?, ?, ?, ?, ?, ?, ?, ?)", (file_id, document_id, final_document_id, promotion_key, outcome, reason, correlation_id, json.dumps(payload or {}, ensure_ascii=False)))
            conn.commit()

    def record_error(self, correlation_id: str, error_code: str, error_message: str, *, file_id: int | None = None, document_id: int | None = None, attempt_id: int | None = None, payload: dict[str, Any] | None = None) -> None:
        with self.work_connection() as conn:
            conn.execute('''insert into errors (correlation_id, file_id, document_id, attempt_id, error_code, error_message, payload_json) values (?, ?, ?, ?, ?, ?, ?)''', (correlation_id, file_id, document_id, attempt_id, error_code, error_message, json.dumps(payload or {}, ensure_ascii=False)))
            conn.execute("insert into system_state(key, value, updated_at) values('last_error', ?, ?) on conflict(key) do update set value=excluded.value, updated_at=excluded.updated_at", (error_message, utc_now()))
            conn.commit()
        self.log_event('error', 'error.recorded', error_message, correlation_id, file_id=file_id, document_id=document_id, attempt_id=attempt_id, payload=payload)

    def set_last_success(self, message: str) -> None:
        with self.work_connection() as conn:
            conn.execute("insert into system_state(key, value, updated_at) values('last_success', ?, ?) on conflict(key) do update set value=excluded.value, updated_at=excluded.updated_at", (message, utc_now()))
            conn.commit()

    def record_ares_validation(self, ico: str, status: str, payload: dict[str, Any]) -> None:
        excerpt = json.dumps(payload, ensure_ascii=False)[:2000]
        with self.work_connection() as conn:
            conn.execute('insert into ares_validations (supplier_ico, checked_at, status, payload_excerpt) values (?, ?, ?, ?)', (ico, utc_now(), status, excerpt))
            conn.commit()

    def finalize_result(self, file_id: int, document_id: int, final_state: str, reason: str, correlation_id: str, *, output_directory: Path | None = None) -> None:
        final_state = self._review_state(final_state)
        with self.work_connection() as work, self.prod_connection() as prod:
            final_document_id = None
            promotion_key = ''
            if final_state == 'final':
                pd, ps = self._processing_snapshot(work, document_id)
                items = work.execute('select * from processing_items where document_id=? order by id asc', (document_id,)).fetchall()
                if not pd or not ps:
                    raise ValueError('Zdrojový dokument pro promotion neexistuje.')
                promotion_key = self._promotion_key(pd, ps)
                blocked_reason = ''
                if not pd['document_number'] or not pd['issued_at'] or float(pd['total_with_vat'] or 0) <= 0:
                    blocked_reason = 'Doklad není kompletní.'
                elif not items:
                    blocked_reason = 'Doklad nemá skutečné položky.'
                elif not ps['ico'] or ps['ares_status'] != 'verified':
                    blocked_reason = 'Dodavatel není kompletní a ověřený.'
                else:
                    blocked_reason = self._validate_processing_items(items, expected_total=float(pd['total_with_vat'] or 0))
                if not blocked_reason:
                    vat_summary, item_total, total_without_vat = self._vat_summary_from_items(items)
                if blocked_reason:
                    work.execute('update files set status=?, last_error=? where id=?', ('retry_pending', blocked_reason, file_id))
                    work.execute('update processing_documents set status=?, final_state=?, promotion_status=?, promotion_blocked_reason=? where id=?', ('retry_pending', 'retry_pending', 'blocked', blocked_reason, document_id))
                    work.execute('''insert into document_results (file_id, document_id, final_state, reason, decided_at, correlation_id, final_document_id) values (?, ?, ?, ?, ?, ?, ?) on conflict(file_id) do update set document_id=excluded.document_id, final_state=excluded.final_state, reason=excluded.reason, decided_at=excluded.decided_at, correlation_id=excluded.correlation_id, final_document_id=excluded.final_document_id''', (file_id, document_id, 'retry_pending', blocked_reason, utc_now(), correlation_id, None))
                    self._sync_working_document_fts(work, file_id)
                    self._set_queue_size(work)
                    work.commit()
                    self.record_promotion_audit(file_id, document_id, None, promotion_key, 'blocked', blocked_reason, correlation_id)
                    self.log_event('warning', 'document.promotion_blocked', blocked_reason, correlation_id, file_id=file_id, document_id=document_id)
                    return
                existing_doc = prod.execute('select id from final_documents where source_document_id=?', (document_id,)).fetchone()
                if existing_doc:
                    final_document_id = int(existing_doc['id'])
                    work.execute('update files set status=?, last_error=? where id=?', ('final', '', file_id))
                    work.execute('update processing_documents set status=?, final_state=?, final_document_id=?, promotion_status=?, promotion_blocked_reason=?, promoted_at=? where id=?', ('final', 'final', final_document_id, 'promoted', '', utc_now(), document_id))
                    work.execute('''insert into document_results (file_id, document_id, final_state, reason, decided_at, correlation_id, final_document_id) values (?, ?, ?, ?, ?, ?, ?) on conflict(file_id) do update set document_id=excluded.document_id, final_state=excluded.final_state, reason=excluded.reason, decided_at=excluded.decided_at, correlation_id=excluded.correlation_id, final_document_id=excluded.final_document_id''', (file_id, document_id, 'final', 'Promotion already existed.', utc_now(), correlation_id, final_document_id))
                    self._sync_working_document_fts(work, file_id)
                    self._set_queue_size(work)
                    work.commit()
                    self.record_promotion_audit(file_id, document_id, final_document_id, promotion_key, 'idempotent', 'Promotion already existed.', correlation_id)
                    return
                business_key = self._business_document_key(pd, ps)
                conflict = prod.execute('select id from final_documents where promotion_key=? or business_document_key=?', (promotion_key, business_key)).fetchone()
                if conflict:
                    raise ValueError('Kolize business klíče při promotion.')
                supplier_row = prod.execute('select id from final_suppliers where ico=?', (ps['ico'],)).fetchone()
                if supplier_row:
                    supplier_id = int(supplier_row['id'])
                    prod.execute('''update final_suppliers set dic=?, name=?, vat_payer=?, address=?, ares_checked_at=?, ares_payload_json=?, source_supplier_id=? where id=?''', (ps['dic'], ps['name'], ps['vat_payer'], ps['address'], utc_now(), ps['ares_payload_json'], ps['id'], supplier_id))
                else:
                    cur = prod.execute('''insert into final_suppliers (ico, dic, name, vat_payer, address, ares_checked_at, ares_payload_json, source_supplier_id) values (?, ?, ?, ?, ?, ?, ?, ?)''', (ps['ico'], ps['dic'], ps['name'] or 'Neznámý dodavatel', ps['vat_payer'], ps['address'], utc_now(), ps['ares_payload_json'], ps['id']))
                    supplier_id = int(cur.lastrowid or 0)
                original = work.execute('select internal_path from files where id=?', (file_id,)).fetchone()[0]
                cur = prod.execute('''insert into final_documents (source_document_id, source_file_id, source_correlation_id, promotion_key, business_document_key, supplier_id, issued_at, document_number, total_without_vat, total_with_vat, vat_summary, extraction_method, original_file_path) values (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'workflow', ?)''', (document_id, file_id, correlation_id, promotion_key, business_key, supplier_id, pd['issued_at'], pd['document_number'], total_without_vat, item_total, vat_summary, original))
                final_document_id = int(cur.lastrowid or 0)
                self._sync_final_document_fts(prod, final_document_id)
                for item in items:
                    icur = prod.execute('''insert into final_items (document_id, source_item_id, name, quantity, unit_price, total_price, vat_rate) values (?, ?, ?, ?, ?, ?, ?)''', (final_document_id, item['id'], item['name'], item['quantity'], item['unit_price'], item['total_price'], item['vat_rate']))
                    self._sync_final_item_fts(prod, int(icur.lastrowid or 0))
                self._refresh_reporting_cache(prod)
            if final_state in {'quarantine', 'unrecognized'}:
                try:
                    self.move_file_to_bucket(file_id, final_state)
                except Exception:
                    pass
            final_state_to_store = final_state
            reason_to_store = reason
            if final_state == 'final' and final_document_id is not None:
                work.execute('update files set status=?, last_error=? where id=?', ('final', '', file_id))
                work.execute('update processing_documents set status=?, final_state=?, final_document_id=?, quarantine_reason=?, promotion_status=?, promotion_blocked_reason=?, promoted_at=? where id=?', ('final', 'final', final_document_id, '', 'promoted', '', utc_now(), document_id))
                reason_to_store = 'Promotion completed.'
            elif final_state in {'quarantine', 'unrecognized'}:
                work.execute('update files set status=?, last_error=? where id=?', (final_state, reason, file_id))
                work.execute('update processing_documents set status=?, final_state=?, final_document_id=?, quarantine_reason=? where id=?', (final_state, final_state, final_document_id, reason, document_id))
            self._sync_working_document_fts(work, file_id)
            work.execute('''insert into document_results (file_id, document_id, final_state, reason, decided_at, correlation_id, final_document_id) values (?, ?, ?, ?, ?, ?, ?) on conflict(file_id) do update set document_id=excluded.document_id, final_state=excluded.final_state, reason=excluded.reason, decided_at=excluded.decided_at, correlation_id=excluded.correlation_id, final_document_id=excluded.final_document_id''', (file_id, document_id, final_state_to_store, reason_to_store, utc_now(), correlation_id, final_document_id))
            self._set_queue_size(work)
            work.commit()
            prod.commit()
        if final_state == 'final':
            self.record_promotion_audit(file_id, document_id, final_document_id, promotion_key, 'promoted', reason_to_store, correlation_id, {'final_document_id': final_document_id})
        self.set_last_success(f'{final_state.upper()} file_id={file_id}')
        self.log_event('info', 'document.finalized', f'Dokument byl ukončen stavem {final_state}.', correlation_id, file_id=file_id, document_id=document_id, payload={'reason': reason_to_store})
    def save_manual_processing_data(self, file_id: int, *, document_id: int | None = None, supplier_ico: str, supplier_name: str, supplier_dic: str, address: str, vat_payer: bool, document_number: str, issued_at: str, total_with_vat: float, items: list[dict[str, Any]] | None = None) -> None:
        with self.work_connection() as conn:
            row = conn.execute('select id, correlation_id from files where id=?', (file_id,)).fetchone()
            if not row:
                raise ValueError('Soubor nebyl nalezen.')
            if document_id is not None:
                doc = conn.execute('select id from processing_documents where file_id=? and id=?', (file_id, document_id)).fetchone()
            else:
                doc = conn.execute('select id from processing_documents where file_id=? order by coalesce(segment_index, 1) asc, id asc', (file_id,)).fetchone()
            if not doc:
                raise ValueError('Procesní dokument nebyl nalezen.')
            document_id = int(doc['id'])
            normalized_ico = self._normalize_ico(supplier_ico)
            normalized_document_number = str(document_number or '').strip()
            normalized_issued_at = self._normalize_issued_at(issued_at)
            normalized_total_with_vat = round(float(total_with_vat or 0), 2)
            manual_items = items or []
            if len(normalized_ico) != 8:
                raise ValueError('ICO musi mit 8 cislic.')
            if not normalized_document_number:
                raise ValueError('Cislo dokladu je povinne.')
            if normalized_total_with_vat <= 0:
                raise ValueError('Celkem s DPH musi byt kladne cislo.')
            item_error = self._validate_processing_items(manual_items, expected_total=normalized_total_with_vat)
            if item_error:
                raise ValueError(item_error)
            payload = {'mode': 'manual', 'document_number': normalized_document_number, 'supplier_ico': normalized_ico, 'total_with_vat': normalized_total_with_vat}
            conn.execute("update processing_documents set document_number=?, issued_at=?, total_with_vat=?, manual_payload_json=?, status='retry_pending', final_state='retry_pending', promotion_status='', promotion_blocked_reason='' where id=?", (normalized_document_number, normalized_issued_at, normalized_total_with_vat, json.dumps(payload, ensure_ascii=False), document_id))
            existing = conn.execute('select id from processing_suppliers where document_id=?', (document_id,)).fetchone()
            if existing:
                conn.execute("update processing_suppliers set ico=?, name=?, dic=?, vat_payer=?, address=?, ares_status='pending', source_layer='manual', ares_payload_json=? where document_id=?", (normalized_ico, supplier_name, supplier_dic, int(vat_payer), address, json.dumps({'manual': True}, ensure_ascii=False), document_id))
                supplier_id = int(existing['id'])
            else:
                cur = conn.execute("insert into processing_suppliers (document_id, ico, dic, name, vat_payer, address, ares_status, ares_payload_json, source_layer) values (?, ?, ?, ?, ?, ?, 'pending', ?, 'manual')", (document_id, normalized_ico, supplier_dic, supplier_name, int(vat_payer), address, json.dumps({'manual': True}, ensure_ascii=False)))
                supplier_id = int(cur.lastrowid or 0)
            conn.execute('update processing_documents set supplier_id=? where id=?', (supplier_id, document_id))
            conn.execute('delete from processing_items where document_id=?', (document_id,))
            for item in manual_items:
                conn.execute("insert into processing_items (document_id, name, quantity, unit_price, total_price, vat_rate, source_kind) values (?, ?, ?, ?, ?, ?, ?)", (document_id, str(item.get('name', normalized_document_number or 'Rucni polozka')).strip(), float(item.get('quantity') or 1), float(item.get('unit_price') or item.get('total_price') or normalized_total_with_vat), float(item.get('total_price') or normalized_total_with_vat), str(item.get('vat_rate') or '').strip(), str(item.get('source_kind') or 'manual')))
            conn.execute('update files set status=?, last_error=? where id=?', ('retry_pending', '', file_id))
            self._sync_working_document_fts(conn, file_id)
            self._set_queue_size(conn)
            conn.commit()


    def list_attempts(self, *, time_filter: str = '', status_filter: str = '', attempt_type: str = '', result_filter: str = '', search: str = '') -> list[sqlite3.Row]:
        sql = '''
            select ea.id, ea.started_at, ea.finished_at, ea.attempt_type, ea.branch, ea.result, ea.reason,
                   ea.document_state, ea.correlation_id, ea.file_id, ea.document_id,
                   f.original_name, f.status as file_status
            from extraction_attempts ea
            join files f on f.id = ea.file_id
            where 1=1
        '''
        params: list[Any] = []
        if time_filter:
            sql += ' and substr(ea.started_at, 1, 10)=?'
            params.append(time_filter)
        if status_filter and status_filter.lower() != 'vše':
            sql += ' and coalesce(ea.document_state, "")=?'
            params.append(status_filter)
        if attempt_type and attempt_type.lower() != 'vše':
            sql += ' and coalesce(ea.attempt_type, "")=?'
            params.append(attempt_type)
        if result_filter and result_filter.lower() != 'vše':
            sql += ' and coalesce(ea.result, "")=?'
            params.append(result_filter)
        if search:
            fts_query = self._fts_query(search)
            if fts_query:
                sql += ' and exists (select 1 from extraction_attempts_fts where extraction_attempts_fts.rowid=ea.id and extraction_attempts_fts match ?)'
                params.append(fts_query)
            else:
                needle = f'%{search.lower()}%'
                sql += ' and (lower(coalesce(f.original_name, "")) like ? or lower(coalesce(ea.reason, "")) like ? or lower(coalesce(ea.branch, "")) like ? or lower(coalesce(ea.processor_name, "")) like ? or lower(coalesce(ea.correlation_id, "")) like ?)'
                params.extend([needle, needle, needle, needle, needle])
        sql += ' order by ea.id desc'
        with self.work_connection() as conn:
            return conn.execute(sql, params).fetchall()

    def list_attempts_page(
        self,
        *,
        time_filter: str = '',
        status_filter: str = '',
        attempt_type: str = '',
        result_filter: str = '',
        search: str = '',
        page: int = 1,
        page_size: int = 100,
    ) -> dict[str, Any]:
        sql = '''
            from extraction_attempts ea
            join files f on f.id = ea.file_id
            where 1=1
        '''
        params: list[Any] = []
        if time_filter:
            sql += ' and substr(ea.started_at, 1, 10)=?'
            params.append(time_filter)
        if status_filter and status_filter.lower() != 'vše':
            sql += ' and coalesce(ea.document_state, "")=?'
            params.append(status_filter)
        if attempt_type and attempt_type.lower() != 'vše':
            sql += ' and coalesce(ea.attempt_type, "")=?'
            params.append(attempt_type)
        if result_filter and result_filter.lower() != 'vše':
            sql += ' and coalesce(ea.result, "")=?'
            params.append(result_filter)
        if search:
            fts_query = self._fts_query(search)
            if fts_query:
                sql += ' and exists (select 1 from extraction_attempts_fts where extraction_attempts_fts.rowid=ea.id and extraction_attempts_fts match ?)'
                params.append(fts_query)
            else:
                needle = f'%{search.lower()}%'
                sql += ' and (lower(coalesce(f.original_name, "")) like ? or lower(coalesce(ea.reason, "")) like ? or lower(coalesce(ea.branch, "")) like ? or lower(coalesce(ea.processor_name, "")) like ? or lower(coalesce(ea.correlation_id, "")) like ?)'
                params.extend([needle, needle, needle, needle, needle])
        page = max(1, int(page or 1))
        page_size = max(1, min(int(page_size or 100), 500))
        offset = (page - 1) * page_size
        with self.work_connection() as conn:
            total = int(conn.execute(f'select count(*) {sql}', params).fetchone()[0])
            rows = conn.execute(
                f'''
                select ea.id, ea.started_at, ea.finished_at, ea.attempt_type, ea.branch, ea.result, ea.reason,
                       ea.document_state, ea.correlation_id, ea.file_id, ea.document_id,
                       f.original_name, f.status as file_status
                {sql}
                order by ea.id desc
                limit ? offset ?
                ''',
                [*params, page_size, offset],
            ).fetchall()
        return {'rows': rows, 'total_count': total, 'page': page, 'page_size': page_size}

    def get_attempt_detail(self, attempt_id: int) -> sqlite3.Row | None:
        with self.work_connection() as conn:
            return conn.execute('''
                select ea.*, f.original_name, f.status as file_status
                from extraction_attempts ea
                join files f on f.id = ea.file_id
                where ea.id=?
            ''', (attempt_id,)).fetchone()

    def get_processing_document_detail(self, file_id: int, document_id: int | None = None) -> sqlite3.Row | None:
        with self.work_connection() as conn:
            sql = '''
                select f.id as file_id, f.original_name, f.internal_path, f.status, f.last_error, f.correlation_id,
                       pd.id as document_id, pd.document_number, pd.issued_at, pd.total_with_vat,
                       pd.quarantine_reason, pd.manual_payload_json, pd.preview_page,
                       pd.promotion_status, pd.promotion_blocked_reason, pd.promoted_at,
                       pd.segment_index, pd.page_from, pd.page_to, pd.split_status, pd.split_confidence,
                       ps.id as supplier_row_id, ps.ico, ps.name as supplier_name, ps.dic,
                       ps.vat_payer, ps.address, ps.ares_status, ps.ares_payload_json
                from files f
                join processing_documents pd on pd.file_id = f.id
                left join processing_suppliers ps on ps.document_id = pd.id
                where f.id=?
            '''
            params: list[Any] = [file_id]
            if document_id is not None:
                sql += ' and pd.id=?'
                params.append(document_id)
            sql += ' order by coalesce(pd.segment_index, 1) asc, pd.id asc limit 1'
            return conn.execute(sql, params).fetchone()

    def _quarantine_category_case_sql(self) -> str:
        return '''
            case
                when lower(coalesce(f.status, '')) = 'duplicate'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%duplic%'
                    then 'Duplicita'
                when lower(coalesce(f.status, '')) = 'unrecognized'
                    or (
                        lower(coalesce(f.status, '')) = 'quarantine'
                        and (
                            lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%offline i openai cesty selhaly%'
                            or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%openai fallback nedodal validni vysledek%'
                            or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%openai vetev selhala%'
                            or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%nerozpozn%'
                            or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%nekompletni%'
                        )
                    )
                    then 'Nerozpoznané'
                when lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%chybějící identifikace%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%chybejici identifikace%'
                    then 'Chybí identifikace'
                when lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%nesoulad součt%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%nesoulad souct%'
                    then 'Nesedí součty'
                when lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%nesoulad dph%'
                    then 'Nesedí DPH'
                when lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%atypická sazba dph%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%atypicka sazba dph%'
                    then 'Atypická sazba DPH'
                when lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%rozpor plátce%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%rozpor platce%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%platce dph%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%na dokladu je neplatce%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%na dokladu je neplátce%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%neplátce%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%neplatce%'
                    then 'Plátce DPH a na dokladu je neplátce DPH'
                else 'Další blokující stav'
            end
        '''

    def list_quarantine_documents(self, *, reason: str = '', search: str = '', category: str = '') -> list[sqlite3.Row]:
        sql = '''
            select f.id as file_id, f.original_name, f.status, f.last_error, f.correlation_id,
                   pd.id as document_id, pd.document_number, pd.issued_at, pd.total_with_vat,
                   coalesce(pd.quarantine_reason, '') as quarantine_reason,
                   pd.segment_index, pd.page_from, pd.page_to, pd.split_status,
                   coalesce(ps.ico, '') as ico, coalesce(ps.name, '') as supplier_name,
                   ''' + self._quarantine_category_case_sql() + ''' as quarantine_category
            from files f
            join processing_documents pd on pd.file_id = f.id
            left join processing_suppliers ps on ps.document_id = pd.id
            where (lower(coalesce(pd.final_state, pd.status, '')) = 'quarantine' or lower(coalesce(f.status, '')) = 'duplicate')
        '''
        params: list[Any] = []
        if reason:
            sql += ' and lower(coalesce(nullif(pd.quarantine_reason, ""), f.last_error, "")) like ?'
            params.append(f'%{reason.lower()}%')
        if category:
            category = self._normalize_quarantine_category(category)
            sql += ' and lower((' + self._quarantine_category_case_sql() + ')) = ?'
            params.append(category.lower())
        if search:
            fts_query = self._fts_query(search)
            if fts_query:
                sql += ' and exists (select 1 from working_documents_fts where working_documents_fts.rowid=f.id and working_documents_fts match ?)'
                params.append(fts_query)
            else:
                needle = f'%{search.lower()}%'
                sql += ' and (lower(coalesce(f.original_name, "")) like ? or lower(coalesce(pd.document_number, "")) like ? or lower(coalesce(ps.ico, "")) like ? or lower(coalesce(ps.name, "")) like ? or lower(coalesce(nullif(pd.quarantine_reason, ""), f.last_error, "")) like ?)'
                params.extend([needle, needle, needle, needle, needle])
        sql += ' order by f.id desc, pd.id desc'
        with self.work_connection() as conn:
            return conn.execute(sql, params).fetchall()

    def list_quarantine_documents_page(
        self,
        *,
        reason: str = '',
        search: str = '',
        category: str = '',
        page: int = 1,
        page_size: int = 100,
    ) -> dict[str, Any]:
        sql = '''
            from files f
            join processing_documents pd on pd.file_id = f.id
            left join processing_suppliers ps on ps.document_id = pd.id
            where (lower(coalesce(pd.final_state, pd.status, '')) = 'quarantine' or lower(coalesce(f.status, '')) = 'duplicate')
        '''
        params: list[Any] = []
        if reason:
            sql += ' and lower(coalesce(nullif(pd.quarantine_reason, ""), f.last_error, "")) like ?'
            params.append(f'%{reason.lower()}%')
        if category:
            category = self._normalize_quarantine_category(category)
            sql += ' and lower((' + self._quarantine_category_case_sql() + ')) = ?'
            params.append(category.lower())
        if search:
            fts_query = self._fts_query(search)
            if fts_query:
                sql += ' and exists (select 1 from working_documents_fts where working_documents_fts.rowid=f.id and working_documents_fts match ?)'
                params.append(fts_query)
            else:
                needle = f'%{search.lower()}%'
                sql += ' and (lower(coalesce(f.original_name, "")) like ? or lower(coalesce(pd.document_number, "")) like ? or lower(coalesce(ps.ico, "")) like ? or lower(coalesce(ps.name, "")) like ? or lower(coalesce(nullif(pd.quarantine_reason, ""), f.last_error, "")) like ?)'
                params.extend([needle, needle, needle, needle, needle])
        page = max(1, int(page or 1))
        page_size = max(1, min(int(page_size or 100), 500))
        offset = (page - 1) * page_size
        with self.work_connection() as conn:
            total = int(conn.execute(f'select count(*) {sql}', params).fetchone()[0])
            rows = conn.execute(
                f'''
                select f.id as file_id, f.original_name, f.internal_path, f.status, f.last_error, f.correlation_id,
                       pd.id as document_id, pd.document_number, pd.issued_at, pd.total_with_vat,
                       coalesce(nullif(pd.quarantine_reason, ''), f.last_error, '') as quarantine_reason,
                       pd.segment_index, pd.page_from, pd.page_to, pd.split_status,
                       coalesce(ps.ico, '') as ico, coalesce(ps.name, '') as supplier_name,
                       {self._quarantine_category_case_sql()} as quarantine_category
                {sql}
                order by f.id desc, pd.id desc
                limit ? offset ?
                ''',
                [*params, page_size, offset],
            ).fetchall()
        return {'rows': rows, 'total_count': total, 'page': page, 'page_size': page_size}

    def _normalize_quarantine_category(self, category: str) -> str:
        normalized = (category or '').strip()
        aliases = {
            'Nesoulad součtů': 'Nesedí součty',
            'Nesoulad DPH': 'Nesedí DPH',
            'Rozpor plátce/neplátce': 'Plátce DPH a na dokladu je neplátce DPH',
            'Chybějící identifikace': 'Chybí identifikace',
            'Nerozpoznané': 'Nerozpoznané',
            'Duplicita': 'Duplicita',
            'Atypická sazba DPH': 'Atypická sazba DPH',
            'Další blokující stav': 'Další blokující stav',
        }
        return aliases.get(normalized, normalized)

    def _expense_review_label(self, category: str) -> str:
        normalized = self._normalize_quarantine_category(category)
        aliases = {
            'Nesedí součty': 'Nesedí součty',
            'Nesedí DPH': 'Nesedí DPH',
            'Chybí identifikace': 'Chybí identifikace',
            'Nerozpoznané': 'Nerozpoznané',
            'Atypická sazba DPH': 'Atypická sazba DPH',
            'Plátce DPH a na dokladu je neplátce DPH': 'Plátce DPH a na dokladu je neplátce DPH',
            'Další blokující stav': 'Další blokující stav',
        }
        return aliases.get(normalized, normalized)

    def list_quarantine_documents(self, *, reason: str = '', search: str = '', category: str = '') -> list[sqlite3.Row]:
        sql = '''
            select f.id as file_id, f.original_name, f.status, f.last_error, f.correlation_id,
                   pd.id as document_id, pd.document_number, pd.issued_at, pd.total_with_vat,
                   coalesce(pd.quarantine_reason, '') as quarantine_reason,
                   pd.segment_index, pd.page_from, pd.page_to, pd.split_status,
                   coalesce(ps.ico, '') as ico, coalesce(ps.name, '') as supplier_name,
                   ''' + self._quarantine_category_case_sql() + ''' as quarantine_category
            from files f
            join processing_documents pd on pd.file_id = f.id
            left join processing_suppliers ps on ps.document_id = pd.id
            where f.status in ('quarantine', 'duplicate')
        '''
        params: list[Any] = []
        if reason:
            sql += ' and lower(coalesce(nullif(pd.quarantine_reason, ""), f.last_error, "")) like ?'
            params.append(f'%{reason.lower()}%')
        if category:
            category = self._normalize_quarantine_category(category)
            sql += ' and lower((' + self._quarantine_category_case_sql() + ')) = ?'
            params.append(category.lower())
        if search:
            fts_query = self._fts_query(search)
            if fts_query:
                sql += ' and exists (select 1 from working_documents_fts where working_documents_fts.rowid=f.id and working_documents_fts match ?)'
                params.append(fts_query)
            else:
                needle = f'%{search.lower()}%'
                sql += ' and (lower(coalesce(f.original_name, "")) like ? or lower(coalesce(pd.document_number, "")) like ? or lower(coalesce(ps.ico, "")) like ? or lower(coalesce(ps.name, "")) like ? or lower(coalesce(nullif(pd.quarantine_reason, ""), f.last_error, "")) like ?)'
                params.extend([needle, needle, needle, needle, needle])
        sql += ' order by f.id desc'
        with self.work_connection() as conn:
            return conn.execute(sql, params).fetchall()

    def list_quarantine_documents_page(
        self,
        *,
        reason: str = '',
        search: str = '',
        category: str = '',
        page: int = 1,
        page_size: int = 100,
    ) -> dict[str, Any]:
        sql = '''
            from files f
            join processing_documents pd on pd.file_id = f.id
            left join processing_suppliers ps on ps.document_id = pd.id
            where f.status in ('quarantine', 'duplicate')
        '''
        params: list[Any] = []
        if reason:
            sql += ' and lower(coalesce(nullif(pd.quarantine_reason, ""), f.last_error, "")) like ?'
            params.append(f'%{reason.lower()}%')
        if category:
            category = self._normalize_quarantine_category(category)
            sql += ' and lower((' + self._quarantine_category_case_sql() + ')) = ?'
            params.append(category.lower())
        if search:
            fts_query = self._fts_query(search)
            if fts_query:
                sql += ' and exists (select 1 from working_documents_fts where working_documents_fts.rowid=f.id and working_documents_fts match ?)'
                params.append(fts_query)
            else:
                needle = f'%{search.lower()}%'
                sql += ' and (lower(coalesce(f.original_name, "")) like ? or lower(coalesce(pd.document_number, "")) like ? or lower(coalesce(ps.ico, "")) like ? or lower(coalesce(ps.name, "")) like ? or lower(coalesce(nullif(pd.quarantine_reason, ""), f.last_error, "")) like ?)'
                params.extend([needle, needle, needle, needle, needle])
        page = max(1, int(page or 1))
        page_size = max(1, min(int(page_size or 100), 500))
        offset = (page - 1) * page_size
        with self.work_connection() as conn:
            total = int(conn.execute(f'select count(*) {sql}', params).fetchone()[0])
            rows = conn.execute(
                f'''
                select f.id as file_id, f.original_name, f.internal_path, f.status, f.last_error, f.correlation_id,
                       pd.id as document_id, pd.document_number, pd.issued_at, pd.total_with_vat,
                       coalesce(nullif(pd.quarantine_reason, ''), f.last_error, '') as quarantine_reason,
                       pd.segment_index, pd.page_from, pd.page_to, pd.split_status,
                       coalesce(ps.ico, '') as ico, coalesce(ps.name, '') as supplier_name,
                       {self._quarantine_category_case_sql()} as quarantine_category
                {sql}
                order by f.id desc
                limit ? offset ?
                ''',
                [*params, page_size, offset],
            ).fetchall()
        return {'rows': rows, 'total_count': total, 'page': page, 'page_size': page_size}

    def list_unrecognized_documents(self, *, search: str = '') -> list[sqlite3.Row]:
        return self.list_quarantine_documents(search=search, category='Nerozpoznané')

    def list_unrecognized_documents_page(
        self,
        *,
        search: str = '',
        page: int = 1,
        page_size: int = 100,
    ) -> dict[str, Any]:
        return self.list_quarantine_documents_page(search=search, category='Nerozpoznané', page=page, page_size=page_size)

    def bulk_mark(self, file_ids: list[int], target_state: str, note: str) -> int:
        if target_state not in {'quarantine', 'unrecognized'}:
            raise ValueError('Nepodporovaný cílový stav.')
        target_state = self._review_state(target_state)
        changed = 0
        with self.work_connection() as conn:
            for file_id in file_ids:
                doc = conn.execute('select id from processing_documents where file_id=?', (file_id,)).fetchone()
                if not doc:
                    continue
                conn.execute('update files set status=?, last_error=? where id=?', (target_state, note, file_id))
                conn.execute('update processing_documents set status=?, final_state=?, quarantine_reason=? where file_id=?', (target_state, target_state, note, file_id))
                conn.execute('''
                    insert into document_results (file_id, document_id, final_state, reason, decided_at, correlation_id, final_document_id)
                    values (?, ?, ?, ?, ?, (select correlation_id from files where id=?), null)
                    on conflict(file_id) do update set
                        document_id=excluded.document_id,
                        final_state=excluded.final_state,
                        reason=excluded.reason,
                        decided_at=excluded.decided_at,
                        correlation_id=excluded.correlation_id,
                        final_document_id=excluded.final_document_id
                ''', (file_id, int(doc['id']), target_state, note, utc_now(), file_id))
                self._sync_working_document_fts(conn, file_id)
                changed += 1
            self._set_queue_size(conn)
            conn.commit()
        return changed

    def mark_retry_pending(self, file_ids: list[int], attempt_type: str = 'manual') -> int:
        changed = 0
        with self.work_connection() as conn:
            for file_id in file_ids:
                doc = conn.execute('select id from processing_documents where file_id=?', (file_id,)).fetchone()
                if not doc:
                    continue
                reason = 'Čeká na ruční pokus.' if attempt_type == 'manual' else 'Čeká na nový pokus.'
                conn.execute('update files set status=?, last_error=? where id=?', ('retry_pending', '', file_id))
                conn.execute("update processing_documents set status='retry_pending', final_state='retry_pending', promotion_status='', promotion_blocked_reason='' where file_id=?", (file_id,))
                conn.execute('''
                    insert into document_results (file_id, document_id, final_state, reason, decided_at, correlation_id, final_document_id)
                    values (?, ?, 'retry_pending', ?, ?, (select correlation_id from files where id=?), null)
                    on conflict(file_id) do update set
                        document_id=excluded.document_id,
                        final_state=excluded.final_state,
                        reason=excluded.reason,
                        decided_at=excluded.decided_at,
                        correlation_id=excluded.correlation_id,
                        final_document_id=excluded.final_document_id
                ''', (file_id, int(doc['id']), reason, utc_now(), file_id))
                self._sync_working_document_fts(conn, file_id)
                changed += 1
            self._set_queue_size(conn)
            conn.commit()
        return changed

    def move_file_to_bucket(self, file_id: int, bucket: str) -> Path:
        mapping = {'quarantine': 'documents/quarantine', 'unrecognized': 'documents/unrecognized'}
        if bucket not in mapping:
            raise ValueError('Neznámý cílový bucket.')
        with self.work_connection() as conn:
            row = conn.execute('select internal_path from files where id=?', (file_id,)).fetchone()
            if not row:
                raise ValueError('Soubor nebyl nalezen.')
            current = self.project_path / row['internal_path']
            target = self.project_path / mapping[bucket] / current.name
            if current.exists() and current.resolve() != target.resolve():
                target.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(current, target)
                conn.execute('update files set internal_path=? where id=?', (str(target.relative_to(self.project_path)), file_id))
                conn.commit()
            return target

    def _safe_output_target(self, base_dir: Path, desired_name: str) -> Path:
        base_dir.mkdir(parents=True, exist_ok=True)
        safe_name = self._safe_file_name(desired_name) or f'file_{uuid.uuid4().hex}'
        candidate = base_dir / safe_name
        if not candidate.exists():
            return candidate
        stem = candidate.stem
        suffix = candidate.suffix
        counter = 1
        while True:
            candidate = base_dir / f'{stem}_{counter}{suffix}'
            if not candidate.exists():
                return candidate
            counter += 1

    def _remove_consumed_input(self, source_path: Path) -> None:
        try:
            if source_path.exists():
                source_path.unlink()
        except FileNotFoundError:
            return
        except Exception:
            self.log_runtime(f'WARN input-cleanup failed path={source_path}')

    def _output_bucket_dir(self, output_directory: Path, bucket: str) -> Path:
        output_root = Path(output_directory).expanduser()
        output_root.mkdir(parents=True, exist_ok=True)
        quarantine_dir = output_root / 'KARANTENA'
        unrecognized_dir = output_root / 'NEROZPOZNANE'
        quarantine_dir.mkdir(parents=True, exist_ok=True)
        unrecognized_dir.mkdir(parents=True, exist_ok=True)
        if bucket == 'final':
            return output_root
        if bucket in {'quarantine', 'duplicate'}:
            return quarantine_dir
        if bucket == 'unrecognized':
            return unrecognized_dir
        raise ValueError('Neznámý output bucket.')

    def _copy_external_output(self, source_path: Path, output_directory: Path, *, bucket: str) -> Path:
        target_dir = self._output_bucket_dir(output_directory, bucket)
        target = self._safe_output_target(target_dir, source_path.name)
        shutil.copy2(source_path, target)
        return target

    def export_processing_document_to_output(self, file_id: int, document_id: int, output_directory: Path, *, bucket: str) -> Path:
        with self.work_connection() as conn:
            row = conn.execute(
                '''
                select f.original_name, f.internal_path, pd.segment_index, pd.page_from, pd.page_to
                from files f
                join processing_documents pd on pd.file_id = f.id
                where f.id=? and pd.id=?
                ''',
                (file_id, document_id),
            ).fetchone()
            document_count = int(conn.execute('select count(*) from processing_documents where file_id=?', (file_id,)).fetchone()[0])
        if not row:
            raise ValueError('Procesní dokument nebyl nalezen.')
        if document_count <= 1:
            return self.export_file_to_output(file_id, output_directory, bucket=bucket)
        source_path = self.project_path / row['internal_path']
        original_name = str(row['original_name'] or source_path.name)
        original_path = Path(original_name)
        desired_name = f'{original_path.stem}__doklad_{int(row["segment_index"] or 1):02d}{original_path.suffix}'
        target_dir = self._output_bucket_dir(output_directory, bucket)
        target = self._safe_output_target(target_dir, desired_name)
        page_from = int(row['page_from'] or 1)
        page_to = int(row['page_to'] or page_from)
        if source_path.suffix.lower() == '.pdf':
            try:
                from pypdf import PdfReader, PdfWriter
            except Exception as exc:
                raise RuntimeError(f'PDF export segmentu není dostupný: {exc}') from exc
            reader = PdfReader(str(source_path))
            writer = PdfWriter()
            for page_no in range(page_from, min(page_to, len(reader.pages)) + 1):
                writer.add_page(reader.pages[page_no - 1])
            with target.open('wb') as handle:
                writer.write(handle)
            return target
        loaded = self.document_loader.load(source_path)
        selected_pages = [page.text for page in loaded.pages if page_from <= int(page.page_no) <= page_to]
        target.write_text('\f'.join(selected_pages), encoding='utf-8')
        return target

    def export_file_to_output(self, file_id: int, output_directory: Path, *, bucket: str) -> Path:
        with self.work_connection() as conn:
            row = conn.execute('select original_name, internal_path from files where id=?', (file_id,)).fetchone()
            if not row:
                raise ValueError('Soubor nebyl nalezen.')
            source_path = self.project_path / row['internal_path']
            desired_name = str(row['original_name'] or source_path.name)
        target_dir = self._output_bucket_dir(output_directory, bucket)
        target = self._safe_output_target(target_dir, desired_name)
        shutil.copy2(source_path, target)
        return target

    def list_visual_patterns(self) -> list[sqlite3.Row]:
        with self.work_connection() as conn:
            return conn.execute('select * from visual_patterns order by is_active desc, name collate nocase asc').fetchall()

    def list_visual_patterns_page(self, *, search: str = '', page: int = 1, page_size: int = 100) -> dict[str, Any]:
        sql = 'from visual_patterns where 1=1'
        params: list[Any] = []
        if search:
            needle = f'%{search.lower()}%'
            sql += ' and (lower(coalesce(name, "")) like ? or lower(coalesce(document_path, "")) like ?)'
            params.extend([needle, needle])
        page = max(1, int(page or 1))
        page_size = max(1, min(int(page_size or 100), 500))
        offset = (page - 1) * page_size
        with self.work_connection() as conn:
            total = int(conn.execute(f'select count(*) {sql}', params).fetchone()[0])
            rows = conn.execute(
                f'''
                select *
                {sql}
                order by is_active desc, name collate nocase asc
                limit ? offset ?
                ''',
                [*params, page_size, offset],
            ).fetchall()
        return {'rows': rows, 'total_count': total, 'page': page, 'page_size': page_size}

    def create_visual_pattern(self, *, name: str, document_path: str, page_no: int, recognition_rules: dict[str, Any], field_map: dict[str, Any], is_active: bool = True, preview_state: dict[str, Any] | None = None) -> int:
        with self.work_connection() as conn:
            cur = conn.execute('insert into visual_patterns(name, is_active, document_path, page_no, recognition_rules, field_map, preview_state, created_at, updated_at) values (?, ?, ?, ?, ?, ?, ?, ?, ?)', (name, int(is_active), document_path, page_no, json.dumps(recognition_rules, ensure_ascii=False), json.dumps(field_map, ensure_ascii=False), json.dumps(preview_state or {}, ensure_ascii=False), utc_now(), utc_now()))
            conn.commit()
            return int(cur.lastrowid)

    def update_visual_pattern(self, pattern_id: int, *, name: str, document_path: str, page_no: int, recognition_rules: dict[str, Any], field_map: dict[str, Any], is_active: bool, preview_state: dict[str, Any] | None = None) -> None:
        with self.work_connection() as conn:
            conn.execute('update visual_patterns set name=?, is_active=?, document_path=?, page_no=?, recognition_rules=?, field_map=?, preview_state=?, updated_at=? where id=?', (name, int(is_active), document_path, page_no, json.dumps(recognition_rules, ensure_ascii=False), json.dumps(field_map, ensure_ascii=False), json.dumps(preview_state or {}, ensure_ascii=False), utc_now(), pattern_id))
            conn.commit()

    def delete_visual_pattern(self, pattern_id: int) -> None:
        with self.work_connection() as conn:
            conn.execute('delete from visual_patterns where id=?', (pattern_id,))
            conn.commit()

    def export_diagnostics_bundle(self, destination: Path) -> Path:
        destination.parent.mkdir(parents=True, exist_ok=True)
        import zipfile
        with zipfile.ZipFile(destination, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
            for src, arcname in [
                (self.project_path / 'logs' / 'runtime.log', 'logs/runtime.log'),
                (self.project_path / 'logs' / 'decisions.jsonl', 'logs/decisions.jsonl'),
                (self.work_db_path, 'work.sqlite3'),
                (self.prod_db_path, 'production.sqlite3'),
                (self.project_path / PROJECT_MARKER, PROJECT_MARKER),
            ]:
                if src.exists():
                    zf.write(src, arcname=arcname)
        return destination

    def create_backup(self) -> Path:
        backup_name = f'kajovospend_backup_{datetime.now(UTC).strftime("%Y%m%d%H%M%S")}.zip'
        return self.export_diagnostics_bundle(self.project_path / 'backups' / backup_name)

    def restore_backup(self, backup_zip: Path) -> None:
        import zipfile
        if not backup_zip.exists():
            raise ValueError('Záloha neexistuje.')
        target_map = {
            'work.sqlite3': self.work_db_path,
            'production.sqlite3': self.prod_db_path,
            PROJECT_MARKER: self.project_path / PROJECT_MARKER,
        }
        with zipfile.ZipFile(backup_zip) as zf:
            for member, target in target_map.items():
                if member in zf.namelist():
                    target.parent.mkdir(parents=True, exist_ok=True)
                    with zf.open(member) as src, open(target, 'wb') as dst:
                        shutil.copyfileobj(src, dst)

    def reset_data_area(self, area: str) -> int:
        area = area.strip().lower()
        with self.work_connection() as work, self.prod_connection() as prod:
            if area == 'processing':
                count = work.execute('select count(*) from files').fetchone()[0]
                for table in ['page_text_blocks', 'document_pages', 'field_candidates', 'line_item_candidates', 'files', 'processing_documents', 'processing_items', 'processing_suppliers', 'extraction_attempts', 'errors', 'document_results', 'ares_validations', 'audit_changes', 'operational_log', 'working_documents_fts', 'extraction_attempts_fts']:
                    try:
                        work.execute(f'delete from {table}')
                    except sqlite3.OperationalError:
                        pass
                work.execute("delete from system_state where key in ('queue_size','last_success','last_error')")
                work.commit()
                return int(count)
            if area == 'production':
                count = prod.execute('select count(*) from final_documents').fetchone()[0]
                for table in ['final_items', 'final_documents', 'final_suppliers', 'final_documents_fts', 'final_items_fts']:
                    try:
                        prod.execute(f'delete from {table}')
                    except sqlite3.OperationalError:
                        pass
                self._refresh_reporting_cache(prod)
                prod.commit()
                return int(count)
            if area == 'patterns':
                count = work.execute('select count(*) from visual_patterns').fetchone()[0]
                work.execute('delete from visual_patterns')
                work.commit()
                return int(count)
            raise ValueError('Neznámá datová oblast pro reset.')

    def list_item_groups(self) -> list[sqlite3.Row]:
        with self.work_connection() as conn:
            return conn.execute('select * from item_groups order by is_active desc, name collate nocase asc').fetchall()

    def list_item_groups_page(self, *, search: str = '', page: int = 1, page_size: int = 100) -> dict[str, Any]:
        sql = 'from item_groups where 1=1'
        params: list[Any] = []
        if search:
            needle = f'%{search.lower()}%'
            sql += ' and (lower(coalesce(name, "")) like ? or lower(coalesce(description, "")) like ?)'
            params.extend([needle, needle])
        page = max(1, int(page or 1))
        page_size = max(1, min(int(page_size or 100), 500))
        offset = (page - 1) * page_size
        with self.work_connection() as conn:
            total = int(conn.execute(f'select count(*) {sql}', params).fetchone()[0])
            rows = conn.execute(
                f'''
                select *
                {sql}
                order by is_active desc, name collate nocase asc
                limit ? offset ?
                ''',
                [*params, page_size, offset],
            ).fetchall()
        return {'rows': rows, 'total_count': total, 'page': page, 'page_size': page_size}

    def create_item_group(self, name: str, description: str = '', is_active: bool = True) -> int:
        with self.work_connection() as conn:
            cur = conn.execute('insert into item_groups(name, description, is_active, created_at, updated_at) values (?, ?, ?, ?, ?)', (name.strip(), description.strip(), int(is_active), utc_now(), utc_now()))
            conn.commit()
            return int(cur.lastrowid)

    def update_item_group(self, group_id: int, name: str, description: str = '', is_active: bool = True) -> None:
        with self.work_connection() as conn:
            conn.execute('update item_groups set name=?, description=?, is_active=?, updated_at=? where id=?', (name.strip(), description.strip(), int(is_active), utc_now(), group_id))
            conn.commit()

    def delete_item_group(self, group_id: int) -> None:
        with self.work_connection() as conn:
            conn.execute('update item_catalog set group_id=null where group_id=?', (group_id,))
            conn.execute('delete from item_groups where id=?', (group_id,))
            conn.commit()

    def list_item_catalog(self) -> list[sqlite3.Row]:
        with self.work_connection() as conn:
            return conn.execute('''select c.*, g.name as group_name from item_catalog c left join item_groups g on g.id=c.group_id order by c.is_active desc, c.name collate nocase asc''').fetchall()

    def list_item_catalog_page(self, *, search: str = '', page: int = 1, page_size: int = 100) -> dict[str, Any]:
        sql = '''
            from item_catalog c
            left join item_groups g on g.id=c.group_id
            where 1=1
        '''
        params: list[Any] = []
        if search:
            needle = f'%{search.lower()}%'
            sql += ' and (lower(coalesce(c.name, "")) like ? or lower(coalesce(c.vat_rate, "")) like ? or lower(coalesce(g.name, "")) like ? or lower(coalesce(c.notes, "")) like ?)'
            params.extend([needle, needle, needle, needle])
        page = max(1, int(page or 1))
        page_size = max(1, min(int(page_size or 100), 500))
        offset = (page - 1) * page_size
        with self.work_connection() as conn:
            total = int(conn.execute(f'select count(*) {sql}', params).fetchone()[0])
            rows = conn.execute(
                f'''
                select c.*, g.name as group_name
                {sql}
                order by c.is_active desc, c.name collate nocase asc
                limit ? offset ?
                ''',
                [*params, page_size, offset],
            ).fetchall()
        return {'rows': rows, 'total_count': total, 'page': page, 'page_size': page_size}

    def create_item_catalog_entry(self, name: str, vat_rate: str = '', group_id: int | None = None, notes: str = '', is_active: bool = True) -> int:
        with self.work_connection() as conn:
            cur = conn.execute('insert into item_catalog(name, vat_rate, group_id, notes, is_active, created_at, updated_at) values (?, ?, ?, ?, ?, ?, ?)', (name.strip(), vat_rate.strip(), group_id, notes.strip(), int(is_active), utc_now(), utc_now()))
            conn.commit()
            return int(cur.lastrowid)

    def update_item_catalog_entry(self, entry_id: int, name: str, vat_rate: str = '', group_id: int | None = None, notes: str = '', is_active: bool = True) -> None:
        with self.work_connection() as conn:
            conn.execute('update item_catalog set name=?, vat_rate=?, group_id=?, notes=?, is_active=?, updated_at=? where id=?', (name.strip(), vat_rate.strip(), group_id, notes.strip(), int(is_active), utc_now(), entry_id))
            conn.commit()

    def delete_item_catalog_entry(self, entry_id: int) -> None:
        with self.work_connection() as conn:
            conn.execute('delete from item_catalog where id=?', (entry_id,))
            conn.commit()

    def list_documents(self) -> list[sqlite3.Row]:
        with self.work_connection() as conn:
            return conn.execute(
                '''
                select f.id as file_id, f.original_name, f.status, f.correlation_id, pd.id as document_id,
                       pd.segment_index, pd.page_from, pd.page_to, pd.split_status,
                       coalesce(pd.final_state, 'pending') as final_state,
                       coalesce(nullif(pd.quarantine_reason, ''), f.last_error, '') as reason
                from files f
                join processing_documents pd on pd.file_id=f.id
                order by f.id desc, coalesce(pd.segment_index, 1) asc, pd.id asc
                '''
            ).fetchall()

    def list_documents_page(self, *, page: int = 1, page_size: int = 100) -> dict[str, Any]:
        page = max(1, int(page or 1))
        page_size = max(1, min(int(page_size or 100), 500))
        offset = (page - 1) * page_size
        sql = '''
            from files f
            join processing_documents pd on pd.file_id=f.id
        '''
        with self.work_connection() as conn:
            total = int(conn.execute(f'select count(*) {sql}').fetchone()[0])
            rows = conn.execute(
                f'''
                select f.id as file_id, f.original_name, f.status, f.correlation_id,
                       pd.id as document_id, pd.segment_index, pd.page_from, pd.page_to, pd.split_status,
                       coalesce(pd.final_state, 'pending') as final_state,
                       coalesce(nullif(pd.quarantine_reason, ''), f.last_error, '') as reason
                {sql}
                order by f.id desc, coalesce(pd.segment_index, 1) asc, pd.id asc
                limit ? offset ?
                ''',
                (page_size, offset),
            ).fetchall()
        return {'rows': rows, 'total_count': total, 'page': page, 'page_size': page_size}

    def separation_report(self) -> dict[str, Any]:
        with self.work_connection() as work, self.prod_connection() as prod:
            work_tables = {row[0] for row in work.execute("select name from sqlite_master where type='table'").fetchall()}
            prod_tables = {row[0] for row in prod.execute("select name from sqlite_master where type='table'").fetchall()}
        workflow_only = {'files', 'processing_documents', 'processing_items', 'processing_suppliers', 'document_pages', 'page_text_blocks', 'field_candidates', 'line_item_candidates', 'extraction_attempts', 'ares_validations'}
        business_only = {'final_documents', 'final_items', 'final_suppliers'}
        same_path = self.work_db_path == self.prod_db_path
        leakage_to_prod = sorted(workflow_only & prod_tables)
        leakage_to_work = sorted(business_only & work_tables)
        ok = not same_path and not leakage_to_prod and not leakage_to_work and (self.working_factory is not self.production_factory)
        return {
            'ok': ok,
            'message': 'Databáze jsou fyzicky i logicky oddělené.' if ok else 'Separace databází selhala.',
            'working_db': str(self.work_db_path),
            'production_db': str(self.prod_db_path),
            'same_canonical_path': same_path,
            'distinct_factories': self.working_factory is not self.production_factory,
            'workflow_tables_in_production': leakage_to_prod,
            'business_tables_in_working': leakage_to_work,
        }

    def operational_panel_data(self) -> dict[str, Any]:
        with self.work_connection() as work:
            recent = [dict(r) for r in work.execute('select original_name, status, last_error from files order by id desc limit 10').fetchall()]
            latest_batch = work.execute(
                '''
                with latest as (
                    select correlation_id
                    from extraction_attempts
                    where coalesce(correlation_id, '') <> ''
                    order by coalesce(finished_at, started_at) desc, id desc
                    limit 1
                )
                select
                    ea.correlation_id,
                    min(ea.started_at) as started_at,
                    max(coalesce(ea.finished_at, ea.started_at)) as finished_at,
                    count(*) as attempt_count,
                    count(distinct ea.file_id) as file_count,
                    coalesce(sum(case when lower(coalesce(ea.result, '')) = 'failed' then 1 else 0 end), 0) as failed_count,
                    min(coalesce(f.original_name, '')) as sample_name
                from extraction_attempts ea
                left join files f on f.id = ea.file_id
                where ea.correlation_id = (select correlation_id from latest)
                group by ea.correlation_id
                '''
            ).fetchone()
            latest_attempt = work.execute(
                '''
                select
                    ea.correlation_id,
                    ea.started_at,
                    ea.finished_at,
                    ea.attempt_type,
                    ea.result,
                    ea.reason,
                    f.original_name
                from extraction_attempts ea
                left join files f on f.id = ea.file_id
                order by coalesce(ea.finished_at, ea.started_at) desc, ea.id desc
                limit 1
                '''
            ).fetchone()
            period_days = 7
            period_start = (datetime.now(UTC) - timedelta(days=period_days)).replace(microsecond=0).isoformat()
            errors_in_period = int(
                work.execute(
                    '''
                    select count(*)
                    from extraction_attempts
                    where lower(coalesce(result, '')) = 'failed'
                      and coalesce(finished_at, started_at, '') >= ?
                    ''',
                    (period_start,),
                ).fetchone()[0]
            )
            started_at = None
            finished_at = None
            last_run_label = 'Žádný běh'
            last_run_source = 'Bez pokusů'
            last_run_attempts = 0
            last_run_errors = 0
            correlation_id = ''
            if latest_batch:
                started_at = latest_batch['started_at']
                finished_at = latest_batch['finished_at']
                last_run_attempts = int(latest_batch['attempt_count'] or 0)
                last_run_errors = int(latest_batch['failed_count'] or 0)
                correlation_id = str(latest_batch['correlation_id'] or '')
                file_count = int(latest_batch['file_count'] or 0)
                sample_name = str(latest_batch['sample_name'] or '').strip()
                last_run_label = f"{finished_at or started_at or '—'}"
                if file_count > 1:
                    last_run_source = f"Batch: {file_count} souboru, {last_run_attempts} pokusu"
                elif sample_name:
                    last_run_source = sample_name
                else:
                    last_run_source = f"{last_run_attempts} pokusu"
            elif latest_attempt:
                started_at = latest_attempt['started_at']
                finished_at = latest_attempt['finished_at']
                last_run_attempts = 1
                last_run_errors = 1 if str(latest_attempt['result'] or '').lower() == 'failed' else 0
                correlation_id = str(latest_attempt['correlation_id'] or '')
                last_run_label = f"{finished_at or started_at or '—'}"
                last_run_source = str(latest_attempt['original_name'] or latest_attempt['attempt_type'] or 'Poslední pokus')
            duration_seconds = None
            started_dt = _parse_timestamp(started_at)
            finished_dt = _parse_timestamp(finished_at)
            if started_dt and finished_dt:
                duration_seconds = max(0, int((finished_dt - started_dt).total_seconds()))
            return {
                'queue_size': int(work.execute("select count(*) from files where status in ('new','queued','processing','retry_pending')").fetchone()[0]),
                'quarantine': int(work.execute("select count(*) from files where status='quarantine'").fetchone()[0]),
                'unrecognized': int(work.execute("select count(*) from files where status='unrecognized'").fetchone()[0]),
                'duplicates': int(work.execute("select count(*) from files where status='duplicate'").fetchone()[0]),
                'recent': recent,
                'last_run': {
                    'label': last_run_label,
                    'source': last_run_source,
                    'started_at': started_at or '',
                    'finished_at': finished_at or '',
                    'duration_seconds': duration_seconds,
                    'attempt_count': last_run_attempts,
                    'failed_count': last_run_errors,
                    'correlation_id': correlation_id,
                },
                'errors_period': {
                    'days': period_days,
                    'count': errors_in_period,
                },
                'logs': {
                    'runtime_log': str(self.runtime_log_path),
                    'decisions_log': str(self.structured_log_path),
                },
            }
    def status_snapshot(self) -> dict[str, str | int]:
        metrics = self.metrics()
        with self.work_connection() as conn:
            last_success = conn.execute("select value from system_state where key='last_success'").fetchone()
            last_error = conn.execute("select value from system_state where key='last_error'").fetchone()
        return {'queue_size': metrics['queue_size'], 'last_success': last_success['value'] if last_success else '—', 'last_error': last_error['value'] if last_error else '—'}

    def list_suppliers(self) -> list[sqlite3.Row]:
        with self.prod_connection() as conn:
            return conn.execute('''
                select fs.id, fs.ico, fs.name, fs.dic, fs.vat_payer, fs.address,
                       count(fd.id) as document_count,
                       coalesce(sum(fd.total_with_vat), 0) as financial_volume
                from final_suppliers fs
                left join final_documents fd on fd.supplier_id=fs.id
                group by fs.id, fs.ico, fs.name, fs.dic, fs.vat_payer, fs.address
                order by fs.name collate nocase asc
            ''').fetchall()

    def get_supplier_detail(self, supplier_id: int) -> sqlite3.Row | None:
        with self.prod_connection() as conn:
            return conn.execute('''
                select fs.*, count(fd.id) as document_count, coalesce(sum(fd.total_with_vat), 0) as financial_volume
                from final_suppliers fs
                left join final_documents fd on fd.supplier_id=fs.id
                where fs.id=?
                group by fs.id
            ''', (supplier_id,)).fetchone()

    def create_supplier(self, *, ico: str, name: str, dic: str, vat_payer: bool, address: str) -> int:
        with self.prod_connection() as conn:
            cur = conn.execute('insert into final_suppliers (ico, dic, name, vat_payer, address, ares_checked_at, ares_payload_json) values (?, ?, ?, ?, ?, ?, ?)', (ico, dic, name, int(vat_payer), address, '', '{}'))
            self._refresh_reporting_cache(conn)
            conn.commit()
            return int(cur.lastrowid)

    def update_supplier(self, supplier_id: int, *, ico: str, name: str, dic: str, vat_payer: bool, address: str, ares_payload: dict[str, Any] | None = None) -> None:
        with self.prod_connection() as conn:
            if ares_payload is None:
                current = conn.execute('select ares_payload_json from final_suppliers where id=?', (supplier_id,)).fetchone()
                payload_json = current['ares_payload_json'] if current else '{}'
                checked = ''
            else:
                payload_json = json.dumps(ares_payload, ensure_ascii=False)
                checked = utc_now()
            conn.execute('''update final_suppliers set ico=?, dic=?, name=?, vat_payer=?, address=?, ares_checked_at=coalesce(nullif(?, ''), ares_checked_at), ares_payload_json=? where id=?''', (ico, dic, name, int(vat_payer), address, checked, payload_json, supplier_id))
            self._refresh_reporting_cache(conn)
            conn.commit()

    def delete_supplier(self, supplier_id: int) -> None:
        with self.prod_connection() as conn:
            count = conn.execute('select count(*) from final_documents where supplier_id=?', (supplier_id,)).fetchone()[0]
            if count:
                raise ValueError('Dodavatele nelze smazat, protože má navázané doklady.')
            conn.execute('delete from final_suppliers where id=?', (supplier_id,))
            self._refresh_reporting_cache(conn)
            conn.commit()

    def merge_suppliers(self, source_supplier_id: int, target_supplier_id: int) -> None:
        if source_supplier_id == target_supplier_id:
            raise ValueError('Zdrojový a cílový dodavatel musí být odlišní.')
        with self.prod_connection() as conn:
            conn.execute('update final_documents set supplier_id=? where supplier_id=?', (target_supplier_id, source_supplier_id))
            conn.execute('delete from final_suppliers where id=?', (source_supplier_id,))
            self._refresh_reporting_cache(conn)
            conn.commit()

    def list_supplier_documents(self, supplier_id: int) -> list[sqlite3.Row]:
        with self.prod_connection() as conn:
            return conn.execute('select id, issued_at, document_number, total_with_vat from final_documents where supplier_id=? order by issued_at desc, id desc', (supplier_id,)).fetchall()

    def list_supplier_documents_page(self, supplier_id: int, *, page: int = 1, page_size: int = 100) -> dict[str, Any]:
        page = max(1, int(page or 1))
        page_size = max(1, min(int(page_size or 100), 500))
        offset = (page - 1) * page_size
        with self.prod_connection() as conn:
            total = int(conn.execute('select count(*) from final_documents where supplier_id=?', (supplier_id,)).fetchone()[0])
            rows = conn.execute(
                '''
                select id, issued_at, document_number, total_with_vat
                from final_documents
                where supplier_id=?
                order by issued_at desc, id desc
                limit ? offset ?
                ''',
                (supplier_id, page_size, offset),
            ).fetchall()
        return {'rows': rows, 'total_count': total, 'page': page, 'page_size': page_size}

    def list_supplier_items(self, supplier_id: int) -> list[sqlite3.Row]:
        with self.prod_connection() as conn:
            return conn.execute('''select fi.id, fi.name, fi.quantity, fi.total_price, fi.vat_rate from final_items fi join final_documents fd on fd.id=fi.document_id where fd.supplier_id=? order by fi.id desc''', (supplier_id,)).fetchall()

    def list_supplier_items_page(self, supplier_id: int, *, page: int = 1, page_size: int = 100) -> dict[str, Any]:
        page = max(1, int(page or 1))
        page_size = max(1, min(int(page_size or 100), 500))
        offset = (page - 1) * page_size
        with self.prod_connection() as conn:
            total = int(conn.execute('select count(*) from final_items fi join final_documents fd on fd.id=fi.document_id where fd.supplier_id=?', (supplier_id,)).fetchone()[0])
            rows = conn.execute(
                '''
                select fi.id, fi.name, fi.quantity, fi.total_price, fi.vat_rate
                from final_items fi
                join final_documents fd on fd.id=fi.document_id
                where fd.supplier_id=?
                order by fi.id desc
                limit ? offset ?
                ''',
                (supplier_id, page_size, offset),
            ).fetchall()
        return {'rows': rows, 'total_count': total, 'page': page, 'page_size': page_size}


    def _reporting_cache_signature(self, conn: sqlite3.Connection) -> dict[str, int]:
        documents = conn.execute('select count(*) as count, coalesce(max(id), 0) as max_id from final_documents').fetchone()
        items = conn.execute('select count(*) as count, coalesce(max(id), 0) as max_id from final_items').fetchone()
        suppliers = conn.execute('select count(*) as count, coalesce(max(id), 0) as max_id from final_suppliers').fetchone()
        return {
            'documents_count': int(documents['count'] or 0),
            'documents_max_id': int(documents['max_id'] or 0),
            'items_count': int(items['count'] or 0),
            'items_max_id': int(items['max_id'] or 0),
            'suppliers_count': int(suppliers['count'] or 0),
            'suppliers_max_id': int(suppliers['max_id'] or 0),
        }

    def _store_reporting_cache_snapshot(self, conn: sqlite3.Connection, cache_key: str, payload: dict[str, Any]) -> None:
        conn.execute(
            '''
            insert into reporting_cache_snapshots(cache_key, payload_json, updated_at)
            values (?, ?, ?)
            on conflict(cache_key) do update set
                payload_json=excluded.payload_json,
                updated_at=excluded.updated_at
            ''',
            (cache_key, json.dumps(payload, ensure_ascii=False), utc_now()),
        )

    def _refresh_reporting_cache(self, conn: sqlite3.Connection) -> None:
        final_documents = conn.execute('select count(*) from final_documents').fetchone()[0]
        suppliers = conn.execute('select count(*) from final_suppliers').fetchone()[0]
        total_amount = conn.execute('select coalesce(sum(total_with_vat),0) from final_documents').fetchone()[0]
        avg_amount = conn.execute('select coalesce(avg(total_with_vat),0) from final_documents').fetchone()[0]
        monthly = conn.execute("select substr(issued_at,1,7) as period, count(*) as count, coalesce(sum(total_with_vat),0) as amount from final_documents group by substr(issued_at,1,7) order by period desc limit 6").fetchall()
        by_month = conn.execute("select substr(issued_at,1,7) as period, round(sum(total_with_vat),2) as amount from final_documents group by substr(issued_at,1,7) order by period desc").fetchall()
        by_year = conn.execute("select substr(issued_at,1,4) as period, round(sum(total_with_vat),2) as amount from final_documents group by substr(issued_at,1,4) order by period desc").fetchall()
        by_quarter = conn.execute("select substr(issued_at,1,4) || '-Q' || cast(((cast(substr(issued_at,6,2) as integer)+2)/3) as integer) as period, round(sum(total_with_vat),2) as amount from final_documents group by 1 order by 1 desc").fetchall()
        vat = conn.execute("select vat_rate, round(sum(total_price),2) as amount from final_items group by vat_rate order by amount desc").fetchall()
        by_supplier_amount = conn.execute("select fs.name, round(sum(fd.total_with_vat),2) as amount from final_documents fd join final_suppliers fs on fs.id=fd.supplier_id group by fs.id order by amount desc limit 10").fetchall()
        by_supplier_docs = conn.execute("select fs.name, count(fd.id) as count from final_documents fd join final_suppliers fs on fs.id=fd.supplier_id group by fs.id order by count desc limit 10").fetchall()
        self._store_reporting_cache_snapshot(
            conn,
            'dashboard',
            {
                'metrics': {
                    'suppliers': int(suppliers),
                    'final_documents': int(final_documents),
                    'total_amount': float(total_amount or 0),
                    'avg_amount': float(avg_amount or 0),
                    'program_state': 'produkční data',
                },
                'monthly_trend': [dict(r) for r in monthly],
            },
        )
        self._store_reporting_cache_snapshot(
            conn,
            'expenses',
            {
                'by_month': [dict(r) for r in by_month],
                'by_quarter': [dict(r) for r in by_quarter],
                'by_year': [dict(r) for r in by_year],
                'vat_breakdown': [dict(r) for r in vat],
                'top_suppliers_by_amount': [dict(r) for r in by_supplier_amount],
                'top_suppliers_by_count': [dict(r) for r in by_supplier_docs],
                'review_states': [],
                'unrecognized_count': 0,
            },
        )
        self._store_reporting_cache_snapshot(conn, 'meta', self._reporting_cache_signature(conn))

    def _load_reporting_cache_snapshot(self, conn: sqlite3.Connection, cache_key: str) -> dict[str, Any]:
        signature = self._reporting_cache_signature(conn)
        meta_row = conn.execute('select payload_json from reporting_cache_snapshots where cache_key=?', ('meta',)).fetchone()
        cached_meta = json.loads(meta_row['payload_json']) if meta_row and meta_row['payload_json'] else None
        if cached_meta != signature:
            self._refresh_reporting_cache(conn)
            conn.commit()
        row = conn.execute('select payload_json from reporting_cache_snapshots where cache_key=?', (cache_key,)).fetchone()
        if not row:
            self._refresh_reporting_cache(conn)
            conn.commit()
            row = conn.execute('select payload_json from reporting_cache_snapshots where cache_key=?', (cache_key,)).fetchone()
        return json.loads(row['payload_json']) if row and row['payload_json'] else {}

    def dashboard_data(self) -> dict[str, Any]:
        with self.prod_connection() as prod:
            final_documents = prod.execute('select count(*) from final_documents').fetchone()[0]
            suppliers = prod.execute('select count(*) from final_suppliers').fetchone()[0]
            total_amount = prod.execute('select coalesce(sum(total_with_vat),0) from final_documents').fetchone()[0]
            avg_amount = prod.execute('select coalesce(avg(total_with_vat),0) from final_documents').fetchone()[0]
            monthly = prod.execute("select substr(issued_at,1,7) as period, count(*) as count, coalesce(sum(total_with_vat),0) as amount from final_documents group by substr(issued_at,1,7) order by period desc limit 6").fetchall()
        return {
            'metrics': {
                'suppliers': int(suppliers),
                'final_documents': int(final_documents),
                'total_amount': float(total_amount or 0),
                'avg_amount': float(avg_amount or 0),
                'program_state': 'produkční data',
            },
            'monthly_trend': [dict(r) for r in monthly],
        }

    def expense_data(self) -> dict[str, Any]:
        with self.prod_connection() as prod:
            by_month = prod.execute("select substr(issued_at,1,7) as period, round(sum(total_with_vat),2) as amount from final_documents group by substr(issued_at,1,7) order by period desc").fetchall()
            by_year = prod.execute("select substr(issued_at,1,4) as period, round(sum(total_with_vat),2) as amount from final_documents group by substr(issued_at,1,4) order by period desc").fetchall()
            by_quarter = prod.execute("select substr(issued_at,1,4) || '-Q' || cast(((cast(substr(issued_at,6,2) as integer)+2)/3) as integer) as period, round(sum(total_with_vat),2) as amount from final_documents group by 1 order by 1 desc").fetchall()
            vat = prod.execute("select vat_rate, round(sum(total_price),2) as amount from final_items group by vat_rate order by amount desc").fetchall()
            by_supplier_amount = prod.execute("select fs.name, round(sum(fd.total_with_vat),2) as amount from final_documents fd join final_suppliers fs on fs.id=fd.supplier_id group by fs.id order by amount desc limit 10").fetchall()
            by_supplier_docs = prod.execute("select fs.name, count(fd.id) as count from final_documents fd join final_suppliers fs on fs.id=fd.supplier_id group by fs.id order by count desc limit 10").fetchall()
        return {
            'by_month': [dict(r) for r in by_month],
            'by_quarter': [dict(r) for r in by_quarter],
            'by_year': [dict(r) for r in by_year],
            'vat_breakdown': [dict(r) for r in vat],
            'top_suppliers_by_amount': [dict(r) for r in by_supplier_amount],
            'top_suppliers_by_count': [dict(r) for r in by_supplier_docs],
            'review_states': [],
            'unrecognized_count': 0,
        }

    def dashboard_data(self) -> dict[str, Any]:
        with self.prod_connection() as prod:
            data = self._load_reporting_cache_snapshot(prod, 'dashboard')
        metrics = data.setdefault('metrics', {})
        metrics['program_state'] = 'produkční data'
        return data

    def expense_data(self) -> dict[str, Any]:
        with self.prod_connection() as prod:
            return self._load_reporting_cache_snapshot(prod, 'expenses')

    def list_final_documents_page(
        self,
        *,
        search: str = '',
        supplier_id: int | None = None,
        period: str = '',
        vat_rate: str = '',
        page: int = 1,
        page_size: int = 100,
    ) -> dict[str, Any]:
        sql = """
            from final_documents fd
            left join final_suppliers fs on fs.id=fd.supplier_id
            where 1=1
        """
        params: list[Any] = []
        if supplier_id:
            sql += ' and fd.supplier_id=?'; params.append(supplier_id)
        if period:
            sql += ' and substr(fd.issued_at,1,7)=?'; params.append(period)
        if vat_rate:
            sql += ' and exists (select 1 from final_items fi where fi.document_id=fd.id and fi.vat_rate=?)'; params.append(vat_rate)
        if search:
            fts_query = self._fts_query(search)
            like = f'%{search.lower()}%'
            if fts_query:
                sql += ' and (exists (select 1 from final_documents_fts where final_documents_fts.rowid=fd.id and final_documents_fts match ?) or lower(coalesce(fs.name, "")) like ? or lower(coalesce(fs.ico, "")) like ?)'
                params.extend([fts_query, like, like])
            else:
                sql += ' and (lower(coalesce(fs.name, "")) like ? or lower(coalesce(fs.ico, "")) like ?)'
                params.extend([like, like])
        page = max(1, int(page or 1))
        page_size = max(1, min(int(page_size or 100), 500))
        offset = (page - 1) * page_size
        with self.prod_connection() as conn:
            total = int(conn.execute(f'select count(*) {sql}', params).fetchone()[0])
            rows = conn.execute(
                f'''
                select fd.id, fd.issued_at, fd.document_number, fd.total_without_vat, fd.total_with_vat, fd.vat_summary,
                       fd.extraction_method, fd.original_file_path, fs.id as supplier_id, fs.name as supplier_name, fs.ico,
                       (select count(*) from final_items fi where fi.document_id=fd.id) as item_count
                {sql}
                order by fd.issued_at desc, fd.id desc
                limit ? offset ?
                ''',
                [*params, page_size, offset],
            ).fetchall()
        return {'rows': rows, 'total_count': total, 'page': page, 'page_size': page_size}

    def list_final_documents(self, *, search: str = '', supplier_id: int | None = None, period: str = '', vat_rate: str = '') -> list[sqlite3.Row]:
        sql = """
            select fd.id, fd.issued_at, fd.document_number, fd.total_without_vat, fd.total_with_vat, fd.vat_summary,
                   fd.extraction_method, fd.original_file_path, fs.id as supplier_id, fs.name as supplier_name, fs.ico,
                   (select count(*) from final_items fi where fi.document_id=fd.id) as item_count
            from final_documents fd
            left join final_suppliers fs on fs.id=fd.supplier_id
            where 1=1
        """
        params: list[Any] = []
        if supplier_id:
            sql += ' and fd.supplier_id=?'; params.append(supplier_id)
        if period:
            sql += ' and substr(fd.issued_at,1,7)=?'; params.append(period)
        if vat_rate:
            sql += ' and exists (select 1 from final_items fi where fi.document_id=fd.id and fi.vat_rate=?)'; params.append(vat_rate)
        if search:
            fts_query = self._fts_query(search)
            like = f'%{search.lower()}%'
            if fts_query:
                sql += ' and (exists (select 1 from final_documents_fts where final_documents_fts.rowid=fd.id and final_documents_fts match ?) or lower(coalesce(fs.name, "")) like ? or lower(coalesce(fs.ico, "")) like ?)'
                params.extend([fts_query, like, like])
            else:
                sql += ' and (lower(coalesce(fs.name, "")) like ? or lower(coalesce(fs.ico, "")) like ?)'
                params.extend([like, like])
        sql += ' order by fd.issued_at desc, fd.id desc limit 1000'
        with self.prod_connection() as conn:
            return conn.execute(sql, params).fetchall()

    def get_final_document_detail(self, document_id: int) -> sqlite3.Row | None:
        with self.prod_connection() as conn:
            return conn.execute("""
                select fd.*, fs.name as supplier_name, fs.ico, fs.dic, fs.address, fs.vat_payer
                from final_documents fd left join final_suppliers fs on fs.id=fd.supplier_id where fd.id=?
            """, (document_id,)).fetchone()

    def list_document_items(self, document_id: int) -> list[sqlite3.Row]:
        with self.prod_connection() as conn:
            return conn.execute('select id, name, quantity, unit_price, total_price, vat_rate from final_items where document_id=? order by id asc', (document_id,)).fetchall()

    def list_final_items(self, *, search: str = '', supplier_id: int | None = None, period: str = '', vat_rate: str = '') -> list[sqlite3.Row]:
        sql = """
            select fi.id, fi.document_id, fi.name, fi.quantity, fi.unit_price, fi.total_price, fi.vat_rate,
                   fd.document_number, fd.issued_at, fs.id as supplier_id, fs.name as supplier_name
            from final_items fi
            join final_documents fd on fd.id=fi.document_id
            left join final_suppliers fs on fs.id=fd.supplier_id
            where 1=1
        """
        params: list[Any] = []
        if supplier_id:
            sql += ' and fs.id=?'; params.append(supplier_id)
        if period:
            sql += ' and substr(fd.issued_at,1,7)=?'; params.append(period)
        if vat_rate:
            sql += ' and fi.vat_rate=?'; params.append(vat_rate)
        if search:
            fts_query = self._fts_query(search)
            like = f'%{search.lower()}%'
            if fts_query:
                sql += ' and (exists (select 1 from final_items_fts where final_items_fts.rowid=fi.id and final_items_fts match ?) or exists (select 1 from final_documents_fts where final_documents_fts.rowid=fd.id and final_documents_fts match ?) or lower(coalesce(fs.name, "")) like ?)'
                params.extend([fts_query, fts_query, like])
            else:
                sql += ' and lower(coalesce(fs.name, "")) like ?'
                params.append(like)
        sql += ' order by fd.issued_at desc, fi.id desc limit 2000'
        with self.prod_connection() as conn:
            return conn.execute(sql, params).fetchall()

    def list_final_items_page(
        self,
        *,
        search: str = '',
        supplier_id: int | None = None,
        period: str = '',
        vat_rate: str = '',
        page: int = 1,
        page_size: int = 200,
    ) -> dict[str, Any]:
        sql = """
            from final_items fi
            join final_documents fd on fd.id=fi.document_id
            left join final_suppliers fs on fs.id=fd.supplier_id
            where 1=1
        """
        params: list[Any] = []
        if supplier_id:
            sql += ' and fs.id=?'; params.append(supplier_id)
        if period:
            sql += ' and substr(fd.issued_at,1,7)=?'; params.append(period)
        if vat_rate:
            sql += ' and fi.vat_rate=?'; params.append(vat_rate)
        if search:
            fts_query = self._fts_query(search)
            like = f'%{search.lower()}%'
            if fts_query:
                sql += ' and (exists (select 1 from final_items_fts where final_items_fts.rowid=fi.id and final_items_fts match ?) or exists (select 1 from final_documents_fts where final_documents_fts.rowid=fd.id and final_documents_fts match ?) or lower(coalesce(fs.name, "")) like ?)'
                params.extend([fts_query, fts_query, like])
            else:
                sql += ' and lower(coalesce(fs.name, "")) like ?'
                params.append(like)
        page = max(1, int(page or 1))
        page_size = max(1, min(int(page_size or 200), 500))
        offset = (page - 1) * page_size
        with self.prod_connection() as conn:
            total = int(conn.execute(f'select count(*) {sql}', params).fetchone()[0])
            rows = conn.execute(
                f'''
                select fi.id, fi.document_id, fi.name, fi.quantity, fi.unit_price, fi.total_price, fi.vat_rate,
                       fd.document_number, fd.issued_at, fs.id as supplier_id, fs.name as supplier_name
                {sql}
                order by fd.issued_at desc, fi.id desc
                limit ? offset ?
                ''',
                [*params, page_size, offset],
            ).fetchall()
        return {'rows': rows, 'total_count': total, 'page': page, 'page_size': page_size}

    def list_suppliers_page(self, *, search: str = '', page: int = 1, page_size: int = 100) -> dict[str, Any]:
        sql = """
            from final_suppliers fs
            left join final_documents fd on fd.supplier_id=fs.id
            where 1=1
        """
        params: list[Any] = []
        if search:
            like = f'%{search.lower()}%'
            sql += ' and (lower(coalesce(fs.ico, "")) like ? or lower(coalesce(fs.name, "")) like ? or lower(coalesce(fs.dic, "")) like ? or lower(coalesce(fs.address, "")) like ?)'
            params.extend([like, like, like, like])
        page = max(1, int(page or 1))
        page_size = max(1, min(int(page_size or 100), 500))
        offset = (page - 1) * page_size
        group_sql = f'''
            select fs.id, fs.ico, fs.name, fs.dic, fs.vat_payer, fs.address,
                   count(fd.id) as document_count,
                   coalesce(sum(fd.total_with_vat), 0) as financial_volume
            {sql}
            group by fs.id, fs.ico, fs.name, fs.dic, fs.vat_payer, fs.address
        '''
        with self.prod_connection() as conn:
            total = int(conn.execute(f'select count(*) from ({group_sql}) as suppliers_page', params).fetchone()[0])
            rows = conn.execute(
                f'''
                {group_sql}
                order by fs.name collate nocase asc
                limit ? offset ?
                ''',
                [*params, page_size, offset],
            ).fetchall()
        return {'rows': rows, 'total_count': total, 'page': page, 'page_size': page_size}

    def get_final_item_detail(self, item_id: int) -> sqlite3.Row | None:
        with self.prod_connection() as conn:
            return conn.execute("""
                select fi.*, fd.document_number, fd.issued_at, fs.id as supplier_id, fs.name as supplier_name
                from final_items fi join final_documents fd on fd.id=fi.document_id
                left join final_suppliers fs on fs.id=fd.supplier_id where fi.id=?
            """, (item_id,)).fetchone()

    def update_final_document(self, document_id: int, **changes: Any) -> None:
        allowed = {'document_number', 'issued_at', 'total_without_vat', 'total_with_vat', 'vat_summary', 'supplier_id'}
        updates = {k: v for k, v in changes.items() if k in allowed}
        if not updates:
            return
        with self.prod_connection() as prod:
            current = prod.execute('select * from final_documents where id=?', (document_id,)).fetchone()
            if not current:
                raise ValueError('Doklad nebyl nalezen.')
            assignments = ', '.join(f'{k}=?' for k in updates)
            params = list(updates.values()) + [document_id]
            prod.execute(f'update final_documents set {assignments} where id=?', params)
            self._sync_final_document_fts(prod, document_id)
            self._refresh_reporting_cache(prod)
            prod.commit()
        with self.work_connection() as work:
            work.execute("""insert into audit_changes (correlation_id, entity_type, entity_id, action, before_json, after_json, note)
                            values (?, 'final_document', ?, 'manual_edit', ?, ?, ?)""", (
                                uuid.uuid4().hex, document_id, json.dumps(dict(current), ensure_ascii=False), json.dumps(updates, ensure_ascii=False), 'Ruční editace finálních dat'))
            work.commit()

    def export_csv(self, rows: list[dict[str, Any]], destination: Path) -> Path:
        import csv
        fieldnames = sorted({k for row in rows for k in row.keys()}) if rows else ['empty']
        with destination.open('w', encoding='utf-8-sig', newline='') as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            writer.writeheader()
            for row in rows:
                writer.writerow(row)
        return destination

    def export_xlsx(self, rows: list[dict[str, Any]], destination: Path) -> Path:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = 'Export'
        fieldnames = sorted({k for row in rows for k in row.keys()}) if rows else ['empty']
        ws.append(fieldnames)
        for row in rows:
            ws.append([row.get(f, '') for f in fieldnames])
        wb.save(destination)
        return destination

    def _reporting_cache_signature(self, conn: sqlite3.Connection, *, schema: str = 'main') -> dict[str, int]:
        final_documents = self._qualified(schema, 'final_documents')
        final_items = self._qualified(schema, 'final_items')
        final_suppliers = self._qualified(schema, 'final_suppliers')
        documents = conn.execute(f'select count(*) as count, coalesce(max(id), 0) as max_id from {final_documents}').fetchone()
        items = conn.execute(f'select count(*) as count, coalesce(max(id), 0) as max_id from {final_items}').fetchone()
        suppliers = conn.execute(f'select count(*) as count, coalesce(max(id), 0) as max_id from {final_suppliers}').fetchone()
        return {
            'documents_count': int(documents['count'] or 0),
            'documents_max_id': int(documents['max_id'] or 0),
            'items_count': int(items['count'] or 0),
            'items_max_id': int(items['max_id'] or 0),
            'suppliers_count': int(suppliers['count'] or 0),
            'suppliers_max_id': int(suppliers['max_id'] or 0),
        }

    def _store_reporting_cache_snapshot(self, conn: sqlite3.Connection, cache_key: str, payload: dict[str, Any], *, schema: str = 'main') -> None:
        reporting_cache_snapshots = self._qualified(schema, 'reporting_cache_snapshots')
        conn.execute(
            f'''
            insert into {reporting_cache_snapshots}(cache_key, payload_json, updated_at)
            values (?, ?, ?)
            on conflict(cache_key) do update set
                payload_json=excluded.payload_json,
                updated_at=excluded.updated_at
            ''',
            (cache_key, json.dumps(payload, ensure_ascii=False), utc_now()),
        )

    def _refresh_reporting_cache(self, conn: sqlite3.Connection, *, schema: str = 'main') -> None:
        final_documents = self._qualified(schema, 'final_documents')
        final_items = self._qualified(schema, 'final_items')
        final_suppliers = self._qualified(schema, 'final_suppliers')
        final_documents_count = conn.execute(f'select count(*) from {final_documents}').fetchone()[0]
        suppliers = conn.execute(f'select count(*) from {final_suppliers}').fetchone()[0]
        total_amount = conn.execute(f'select coalesce(sum(total_with_vat),0) from {final_documents}').fetchone()[0]
        avg_amount = conn.execute(f'select coalesce(avg(total_with_vat),0) from {final_documents}').fetchone()[0]
        monthly = conn.execute(f"select substr(issued_at,1,7) as period, count(*) as count, coalesce(sum(total_with_vat),0) as amount from {final_documents} group by substr(issued_at,1,7) order by period desc limit 6").fetchall()
        by_month = conn.execute(f"select substr(issued_at,1,7) as period, round(sum(total_with_vat),2) as amount from {final_documents} group by substr(issued_at,1,7) order by period desc").fetchall()
        by_year = conn.execute(f"select substr(issued_at,1,4) as period, round(sum(total_with_vat),2) as amount from {final_documents} group by substr(issued_at,1,4) order by period desc").fetchall()
        by_quarter = conn.execute(f"select substr(issued_at,1,4) || '-Q' || cast(((cast(substr(issued_at,6,2) as integer)+2)/3) as integer) as period, round(sum(total_with_vat),2) as amount from {final_documents} group by 1 order by 1 desc").fetchall()
        vat = conn.execute(f"select vat_rate, round(sum(total_price),2) as amount from {final_items} group by vat_rate order by amount desc").fetchall()
        by_supplier_amount = conn.execute(f"select fs.name, round(sum(fd.total_with_vat),2) as amount from {final_documents} fd join {final_suppliers} fs on fs.id=fd.supplier_id group by fs.id order by amount desc limit 10").fetchall()
        by_supplier_docs = conn.execute(f"select fs.name, count(fd.id) as count from {final_documents} fd join {final_suppliers} fs on fs.id=fd.supplier_id group by fs.id order by count desc limit 10").fetchall()
        self._store_reporting_cache_snapshot(
            conn,
            'dashboard',
            {
                'metrics': {
                    'suppliers': int(suppliers),
                    'final_documents': int(final_documents_count),
                    'total_amount': float(total_amount or 0),
                    'avg_amount': float(avg_amount or 0),
                    'program_state': 'produkční data',
                },
                'monthly_trend': [dict(r) for r in monthly],
            },
            schema=schema,
        )
        self._store_reporting_cache_snapshot(
            conn,
            'expenses',
            {
                'by_month': [dict(r) for r in by_month],
                'by_quarter': [dict(r) for r in by_quarter],
                'by_year': [dict(r) for r in by_year],
                'vat_breakdown': [dict(r) for r in vat],
                'top_suppliers_by_amount': [dict(r) for r in by_supplier_amount],
                'top_suppliers_by_count': [dict(r) for r in by_supplier_docs],
                'review_states': [],
                'unrecognized_count': 0,
            },
            schema=schema,
        )
        self._store_reporting_cache_snapshot(conn, 'meta', self._reporting_cache_signature(conn, schema=schema), schema=schema)

    def _load_reporting_cache_snapshot(self, conn: sqlite3.Connection, cache_key: str, *, schema: str = 'main') -> dict[str, Any]:
        reporting_cache_snapshots = self._qualified(schema, 'reporting_cache_snapshots')
        signature = self._reporting_cache_signature(conn, schema=schema)
        meta_row = conn.execute(f'select payload_json from {reporting_cache_snapshots} where cache_key=?', ('meta',)).fetchone()
        cached_meta = json.loads(meta_row['payload_json']) if meta_row and meta_row['payload_json'] else None
        if cached_meta != signature:
            self._refresh_reporting_cache(conn, schema=schema)
            conn.commit()
        row = conn.execute(f'select payload_json from {reporting_cache_snapshots} where cache_key=?', (cache_key,)).fetchone()
        if not row:
            self._refresh_reporting_cache(conn, schema=schema)
            conn.commit()
            row = conn.execute(f'select payload_json from {reporting_cache_snapshots} where cache_key=?', (cache_key,)).fetchone()
        return json.loads(row['payload_json']) if row and row['payload_json'] else {}

    def _operational_dashboard_metrics(self, conn: sqlite3.Connection) -> dict[str, Any]:
        conn.execute('attach database ? as prod_db', (str(self.prod_db_path),))
        try:
            queue_size = int(conn.execute("select count(*) from files where status in ('new','queued','processing','retry_pending')").fetchone()[0])
            processed_files = int(conn.execute("select count(*) from files where status in ('final','quarantine','unrecognized','duplicate')").fetchone()[0])
            total_files = int(conn.execute('select count(*) from files').fetchone()[0])
            quarantine_files = int(conn.execute("select count(*) from files where status='quarantine'").fetchone()[0])
            unrecognized_files = int(conn.execute("select count(*) from files where status='unrecognized'").fetchone()[0])
            stored_accounts = int(conn.execute('select count(*) from prod_db.final_documents').fetchone()[0])
            breakdown_row = conn.execute(
                '''
                with final_sources as (
                    select fd.id as final_document_id, fd.source_document_id
                    from prod_db.final_documents fd
                ),
                classified as (
                    select
                        fs.final_document_id,
                        case
                            when exists(
                                select 1
                                from processing_items pi
                                where pi.document_id = fs.source_document_id
                                  and lower(coalesce(pi.source_kind, '')) = 'manual'
                            )
                            or exists(
                                select 1
                                from extraction_attempts ea
                                where ea.document_id = fs.source_document_id
                                  and lower(coalesce(ea.attempt_type, '')) = 'manual'
                                  and lower(coalesce(ea.result, '')) = 'success'
                            )
                                then 'manual'
                            when exists(
                                select 1
                                from extraction_attempts ea
                                where ea.document_id = fs.source_document_id
                                  and lower(coalesce(ea.attempt_type, '')) = 'openai'
                                  and lower(coalesce(ea.result, '')) = 'success'
                            )
                                then 'api'
                            else 'offline'
                        end as extraction_origin
                    from final_sources fs
                )
                select
                    coalesce(sum(case when extraction_origin='offline' then 1 else 0 end), 0) as offline_count,
                    coalesce(sum(case when extraction_origin='api' then 1 else 0 end), 0) as api_count,
                    coalesce(sum(case when extraction_origin='manual' then 1 else 0 end), 0) as manual_count
                from classified
                '''
            ).fetchone()
            return {
                'unprocessed_files': queue_size,
                'processed_files': processed_files,
                'stored_accounts': stored_accounts,
                'total_files': total_files,
                'quarantine_files': quarantine_files,
                'unrecognized_files': unrecognized_files,
                'success_breakdown': {
                    'offline': int(breakdown_row['offline_count'] or 0),
                    'api': int(breakdown_row['api_count'] or 0),
                    'manual': int(breakdown_row['manual_count'] or 0),
                },
            }
        finally:
            conn.execute('detach database prod_db')

    def dashboard_data(self) -> dict[str, Any]:
        with self.prod_connection() as prod:
            production = self._load_reporting_cache_snapshot(prod, 'dashboard')
        with self.work_connection() as work:
            operations = self._operational_dashboard_metrics(work)
        metrics = production.setdefault('metrics', {})
        metrics['program_state'] = 'produkční data'
        return {
            'operations': operations,
            'production': production,
            'metrics': metrics,
            'monthly_trend': production.get('monthly_trend', []),
            'success_breakdown': dict(operations['success_breakdown']),
        }

    def _expense_review_metrics(self, conn: sqlite3.Connection) -> dict[str, Any]:
        category_sql = self._quarantine_category_case_sql()
        rows = conn.execute(
            f'''
            select quarantine_category, count(*) as count
            from (
                select {category_sql} as quarantine_category
                from files f
                join processing_documents pd on pd.file_id = f.id
                where f.status in ('quarantine', 'duplicate')
            ) review_categories
            group by quarantine_category
            order by count desc, quarantine_category asc
            '''
        ).fetchall()
        grouped = {str(row['quarantine_category']): int(row['count'] or 0) for row in rows}
        unrecognized_count = int(grouped.pop('Nerozpoznané', 0))
        review_states = [{'status': key, 'count': value} for key, value in grouped.items() if value > 0]
        return {
            'review_states': review_states,
            'unrecognized_count': unrecognized_count,
        }

    def expense_data(self) -> dict[str, Any]:
        with self.prod_connection() as prod:
            expenses = self._load_reporting_cache_snapshot(prod, 'expenses')
        with self.work_connection() as work:
            review_metrics = self._expense_review_metrics(work)
        expenses['review_states'] = review_metrics['review_states']
        expenses['requires_review_total'] = review_metrics['requires_review_total']
        expenses['requires_review_definition'] = review_metrics['requires_review_definition']
        expenses['unrecognized_count'] = review_metrics['unrecognized_count']
        expenses['unrecognized_definition'] = review_metrics['unrecognized_definition']
        return expenses

    def _quarantine_category_case_sql(self) -> str:
        return '''
            case
                when lower(coalesce(f.status, '')) = 'duplicate'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%duplic%'
                    then 'Duplicita'
                when lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%chybějící identifikace%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%chybejici identifikace%'
                    then 'Chybí identifikace'
                when lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%nesoulad součt%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%nesoulad souct%'
                    then 'Nesedí součty'
                when lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%nesoulad dph%'
                    then 'Nesedí DPH'
                when lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%atypická sazba dph%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%atypicka sazba dph%'
                    then 'Atypická sazba DPH'
                when lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%rozpor plátce%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%rozpor platce%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%platce dph%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%na dokladu je neplatce%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%na dokladu je neplátce%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%neplátce%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%neplatce%'
                    then 'Plátce DPH a na dokladu je neplátce DPH'
                else 'Další blokující stav'
            end
        '''

    def list_unrecognized_documents(self, *, search: str = '') -> list[sqlite3.Row]:
        sql = '''
            select f.id as file_id, f.original_name, f.status, f.last_error, f.correlation_id,
                   pd.id as document_id, pd.document_number, pd.issued_at, pd.total_with_vat,
                   coalesce(nullif(pd.quarantine_reason, ''), f.last_error, '') as quarantine_reason,
                   pd.segment_index, pd.page_from, pd.page_to, pd.split_status,
                   coalesce(ps.ico, '') as ico, coalesce(ps.name, '') as supplier_name
            from files f
            join processing_documents pd on pd.file_id = f.id
            left join processing_suppliers ps on ps.document_id = pd.id
            where lower(coalesce(pd.final_state, pd.status, f.status, '')) = 'unrecognized'
        '''
        params: list[Any] = []
        if search:
            fts_query = self._fts_query(search)
            if fts_query:
                sql += ' and exists (select 1 from working_documents_fts where working_documents_fts.rowid=f.id and working_documents_fts match ?)'
                params.append(fts_query)
            else:
                needle = f'%{search.lower()}%'
                sql += ' and (lower(coalesce(f.original_name, "")) like ? or lower(coalesce(pd.document_number, "")) like ? or lower(coalesce(ps.ico, "")) like ? or lower(coalesce(ps.name, "")) like ? or lower(coalesce(nullif(pd.quarantine_reason, ""), f.last_error, "")) like ?)'
                params.extend([needle, needle, needle, needle, needle])
        sql += ' order by f.id desc'
        with self.work_connection() as conn:
            return conn.execute(sql, params).fetchall()

    def list_unrecognized_documents_page(
        self,
        *,
        search: str = '',
        page: int = 1,
        page_size: int = 100,
    ) -> dict[str, Any]:
        sql = '''
            from files f
            join processing_documents pd on pd.file_id = f.id
            left join processing_suppliers ps on ps.document_id = pd.id
            where lower(coalesce(pd.final_state, pd.status, f.status, '')) = 'unrecognized'
        '''
        params: list[Any] = []
        if search:
            fts_query = self._fts_query(search)
            if fts_query:
                sql += ' and exists (select 1 from working_documents_fts where working_documents_fts.rowid=f.id and working_documents_fts match ?)'
                params.append(fts_query)
            else:
                needle = f'%{search.lower()}%'
                sql += ' and (lower(coalesce(f.original_name, "")) like ? or lower(coalesce(pd.document_number, "")) like ? or lower(coalesce(ps.ico, "")) like ? or lower(coalesce(ps.name, "")) like ? or lower(coalesce(nullif(pd.quarantine_reason, ""), f.last_error, "")) like ?)'
                params.extend([needle, needle, needle, needle, needle])
        page = max(1, int(page or 1))
        page_size = max(1, min(int(page_size or 100), 500))
        offset = (page - 1) * page_size
        with self.work_connection() as conn:
            total = int(conn.execute(f'select count(*) {sql}', params).fetchone()[0])
            rows = conn.execute(
                f'''
                select f.id as file_id, f.original_name, f.internal_path, f.status, f.last_error, f.correlation_id,
                       pd.id as document_id, pd.document_number, pd.issued_at, pd.total_with_vat,
                       coalesce(nullif(pd.quarantine_reason, ''), f.last_error, '') as quarantine_reason,
                       pd.segment_index, pd.page_from, pd.page_to, pd.split_status,
                       coalesce(ps.ico, '') as ico, coalesce(ps.name, '') as supplier_name
                {sql}
                order by f.id desc
                limit ? offset ?
                ''',
                [*params, page_size, offset],
            ).fetchall()
        return {'rows': rows, 'total_count': total, 'page': page, 'page_size': page_size}

    def _expense_review_metrics(self, conn: sqlite3.Connection) -> dict[str, Any]:
        category_sql = self._quarantine_category_case_sql()
        rows = conn.execute(
            f'''
            select quarantine_category, count(*) as count
            from (
                select {category_sql} as quarantine_category
                from files f
                join processing_documents pd on pd.file_id = f.id
                where (lower(coalesce(pd.final_state, pd.status, '')) = 'quarantine' or lower(coalesce(f.status, '')) = 'duplicate')
            ) review_categories
            group by quarantine_category
            order by count desc, quarantine_category asc
            '''
        ).fetchall()
        grouped = {str(row['quarantine_category']): int(row['count'] or 0) for row in rows}
        unrecognized_count = int(conn.execute("select count(*) from processing_documents where lower(coalesce(final_state, status, ''))='unrecognized'").fetchone()[0])
        review_states = [{'status': key, 'count': value} for key, value in grouped.items() if value > 0]
        requires_review_total = sum(int(row['count']) for row in review_states)
        return {
            'review_states': review_states,
            'requires_review_total': requires_review_total,
            'requires_review_definition': 'Doklady ve frontě KARANTÉNA a duplicity, které vyžadují kontrolu nebo zásah uživatele.',
            'unrecognized_count': unrecognized_count,
            'unrecognized_definition': 'Doklady ve skutečné frontě NEROZPOZNANÉ, kde se nepodařilo spolehlivě vytěžit obsah.',
        }

    def _quarantine_category_case_sql(self) -> str:
        return '''
            case
                when lower(coalesce(f.status, '')) = 'duplicate'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%duplic%'
                    then 'Duplicita'
                when lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%chyb\u011bj\u00edc\u00ed identifikace%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%chybejici identifikace%'
                    then 'Chyb\u00ed identifikace'
                when lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%nesoulad sou\u010dt%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%nesoulad souct%'
                    then 'Nesed\u00ed sou\u010dty'
                when lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%nesoulad dph%'
                    then 'Nesed\u00ed DPH'
                when lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%atypick\u00e1 sazba dph%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%atypicka sazba dph%'
                    then 'Atypick\u00e1 sazba DPH'
                when lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%rozpor pl\u00e1tce%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%rozpor platce%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%platce dph%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%na dokladu je neplatce%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%na dokladu je nepl\u00e1tce%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%nepl\u00e1tce%'
                    or lower(coalesce(pd.quarantine_reason, f.last_error, '')) like '%neplatce%'
                    then 'Pl\u00e1tce DPH a na dokladu je nepl\u00e1tce DPH'
                else 'Dal\u0161\u00ed blokuj\u00edc\u00ed stav'
            end
        '''

    def export_diagnostics_bundle(self, destination: Path) -> Path:
        destination.parent.mkdir(parents=True, exist_ok=True)
        import zipfile
        import tempfile
        with zipfile.ZipFile(destination, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
            for src, arcname in [
                (self.project_path / 'logs' / 'runtime.log', 'logs/runtime.log'),
                (self.project_path / 'logs' / 'decisions.jsonl', 'logs/decisions.jsonl'),
                (self.project_path / PROJECT_MARKER, PROJECT_MARKER),
            ]:
                if src.exists():
                    zf.write(src, arcname=arcname)
            for source_db, arcname in [
                (self.work_db_path, 'work.sqlite3'),
                (self.prod_db_path, 'production.sqlite3'),
            ]:
                if not source_db.exists():
                    continue
                with tempfile.NamedTemporaryFile(delete=False, suffix='.sqlite3') as handle:
                    snapshot_path = Path(handle.name)
                try:
                    with sqlite3.connect(source_db) as src_conn, sqlite3.connect(snapshot_path) as snapshot_conn:
                        src_conn.backup(snapshot_conn)
                    zf.write(snapshot_path, arcname=arcname)
                finally:
                    try:
                        snapshot_path.unlink(missing_ok=True)
                    except PermissionError:
                        pass
        return destination

    def restore_backup(self, backup_zip: Path) -> None:
        import zipfile
        if not backup_zip.exists():
            raise ValueError('Zaloha neexistuje.')
        target_map = {
            'work.sqlite3': self.work_db_path,
            'production.sqlite3': self.prod_db_path,
            PROJECT_MARKER: self.project_path / PROJECT_MARKER,
        }
        with zipfile.ZipFile(backup_zip) as zf:
            for member, target in target_map.items():
                if member not in zf.namelist():
                    continue
                target.parent.mkdir(parents=True, exist_ok=True)
                payload = zf.read(member)
                if Path(member).suffix == '.sqlite3':
                    restored_target = target.with_name(f'.restore-{target.name}')
                    restored_target.write_bytes(payload)
                    with sqlite3.connect(restored_target) as verify_conn:
                        if str(verify_conn.execute('pragma integrity_check').fetchone()[0]).lower() != 'ok':
                            raise ValueError(f'Obnova databaze {target.name} selhala pri integrity_check.')
                    with sqlite3.connect(restored_target) as source_conn, sqlite3.connect(target) as target_conn:
                        source_conn.backup(target_conn)
                    try:
                        restored_target.unlink(missing_ok=True)
                    except PermissionError:
                        pass
                else:
                    target.write_bytes(payload)
        self.paths = ProjectDatabasePaths.from_project_root(self.project_path)
        self.working_factory = WorkingConnectionFactory(self.paths)
        self.production_factory = ProductionConnectionFactory(self.paths)

    def reset_data_area(self, area: str) -> int:
        area = area.strip().lower()
        with self.work_connection() as work, self.prod_connection() as prod:
            if area == 'processing':
                count = work.execute('select count(*) from files').fetchone()[0]
                for table in ['page_text_blocks', 'field_candidates', 'line_item_candidates', 'document_pages', 'processing_items', 'processing_suppliers', 'extraction_attempts', 'document_results', 'errors', 'ares_validations', 'audit_changes', 'operational_log', 'processing_documents', 'files', 'working_documents_fts', 'extraction_attempts_fts']:
                    try:
                        work.execute(f'delete from {table}')
                    except sqlite3.OperationalError:
                        pass
                work.execute("delete from system_state where key in ('queue_size','last_success','last_error')")
                work.commit()
                return int(count)
            if area == 'production':
                count = prod.execute('select count(*) from final_documents').fetchone()[0]
                for table in ['final_items', 'final_documents', 'final_suppliers', 'final_documents_fts', 'final_items_fts']:
                    try:
                        prod.execute(f'delete from {table}')
                    except sqlite3.OperationalError:
                        pass
                self._refresh_reporting_cache(prod)
                prod.commit()
                return int(count)
            if area == 'patterns':
                count = work.execute('select count(*) from visual_patterns').fetchone()[0]
                work.execute('delete from visual_patterns')
                work.commit()
                return int(count)
            raise ValueError('Neznama datova oblast pro reset.')

    def finalize_result(self, file_id: int, document_id: int, final_state: str, reason: str, correlation_id: str, *, output_directory: Path | None = None) -> None:
        final_state = self._review_state(final_state)
        final_document_id = None
        promotion_key = ''
        audit_outcome = ''
        reason_to_store = reason
        with self.work_connection() as conn:
            conn.execute('attach database ? as prod_db', (str(self.prod_db_path),))
            try:
                if final_state == 'final':
                    pd, ps = self._processing_snapshot(conn, document_id)
                    items = conn.execute('select * from processing_items where document_id=? order by id asc', (document_id,)).fetchall()
                    if not pd or not ps:
                        raise ValueError('Zdrojovy dokument pro promotion neexistuje.')
                    promotion_key = self._promotion_key(pd, ps)
                    blocked_reason = ''
                    if not pd['document_number'] or not pd['issued_at'] or float(pd['total_with_vat'] or 0) <= 0:
                        blocked_reason = 'Doklad neni kompletni.'
                    elif not items:
                        blocked_reason = 'Doklad nema skutecne polozky.'
                    elif not ps['ico'] or ps['ares_status'] != 'verified':
                        blocked_reason = 'Dodavatel neni kompletni a overeny.'
                    else:
                        blocked_reason = self._validate_processing_items(items, expected_total=float(pd['total_with_vat'] or 0))
                    if blocked_reason:
                        conn.execute('update files set status=?, last_error=? where id=?', ('retry_pending', blocked_reason, file_id))
                        conn.execute('update processing_documents set status=?, final_state=?, promotion_status=?, promotion_blocked_reason=? where id=?', ('retry_pending', 'retry_pending', 'blocked', blocked_reason, document_id))
                        conn.execute('''insert into document_results (file_id, document_id, final_state, reason, decided_at, correlation_id, final_document_id) values (?, ?, ?, ?, ?, ?, ?) on conflict(file_id) do update set document_id=excluded.document_id, final_state=excluded.final_state, reason=excluded.reason, decided_at=excluded.decided_at, correlation_id=excluded.correlation_id, final_document_id=excluded.final_document_id''', (file_id, document_id, 'retry_pending', blocked_reason, utc_now(), correlation_id, None))
                        self._sync_working_document_fts(conn, file_id)
                        self._set_queue_size(conn)
                        conn.commit()
                        self.record_promotion_audit(file_id, document_id, None, promotion_key, 'blocked', blocked_reason, correlation_id)
                        self.log_event('warning', 'document.promotion_blocked', blocked_reason, correlation_id, file_id=file_id, document_id=document_id)
                        return
                    vat_summary, item_total, total_without_vat = self._vat_summary_from_items(items)
                    business_key = self._business_document_key(pd, ps)
                    original = conn.execute('select internal_path from files where id=?', (file_id,)).fetchone()[0]
                    existing_doc = conn.execute('select id from prod_db.final_documents where source_document_id=?', (document_id,)).fetchone()
                    if existing_doc:
                        final_document_id = int(existing_doc['id'])
                        conflict = conn.execute('select id from prod_db.final_documents where (promotion_key=? or business_document_key=?) and id<>?', (promotion_key, business_key, final_document_id)).fetchone()
                    else:
                        conflict = conn.execute('select id from prod_db.final_documents where promotion_key=? or business_document_key=?', (promotion_key, business_key)).fetchone()
                    if conflict:
                        raise ValueError('Kolize business klice pri promotion.')
                    supplier_id = self._upsert_final_supplier(conn, ps, schema='prod_db')
                    if existing_doc and self._final_snapshot_matches(conn, final_document_id, supplier_id=supplier_id, pd=pd, items=items, vat_summary=vat_summary, item_total=item_total, total_without_vat=total_without_vat, original_path=original, schema='prod_db'):
                        audit_outcome = 'idempotent'
                        reason_to_store = 'Promotion already existed.'
                    else:
                        if existing_doc:
                            conn.execute('''update prod_db.final_documents set source_file_id=?, source_correlation_id=?, promotion_key=?, business_document_key=?, supplier_id=?, issued_at=?, document_number=?, total_without_vat=?, total_with_vat=?, vat_summary=?, extraction_method='workflow', original_file_path=? where id=?''', (file_id, correlation_id, promotion_key, business_key, supplier_id, pd['issued_at'], pd['document_number'], total_without_vat, item_total, vat_summary, original, final_document_id))
                            audit_outcome = 'reconciled'
                            reason_to_store = 'Promotion reconciled existing production document.'
                        else:
                            cur = conn.execute('''insert into prod_db.final_documents (source_document_id, source_file_id, source_correlation_id, promotion_key, business_document_key, supplier_id, issued_at, document_number, total_without_vat, total_with_vat, vat_summary, extraction_method, original_file_path) values (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'workflow', ?)''', (document_id, file_id, correlation_id, promotion_key, business_key, supplier_id, pd['issued_at'], pd['document_number'], total_without_vat, item_total, vat_summary, original))
                            final_document_id = int(cur.lastrowid or 0)
                            audit_outcome = 'promoted'
                            reason_to_store = 'Promotion completed.'
                        self._sync_final_document_fts(conn, final_document_id, schema='prod_db')
                        self._replace_final_items(conn, final_document_id, items, schema='prod_db')
                        self._refresh_reporting_cache(conn, schema='prod_db')
                elif final_state in {'quarantine', 'unrecognized'}:
                    document_count = int(conn.execute('select count(*) from processing_documents where file_id=?', (file_id,)).fetchone()[0])
                    if document_count <= 1:
                        try:
                            self.move_file_to_bucket(file_id, final_state)
                        except Exception:
                            pass
                if final_state == 'final' and final_document_id is not None:
                    conn.execute('update processing_documents set status=?, final_state=?, final_document_id=?, quarantine_reason=?, promotion_status=?, promotion_blocked_reason=?, promoted_at=? where id=?', ('final', 'final', final_document_id, '', 'promoted', '', utc_now(), document_id))
                elif final_state in {'quarantine', 'unrecognized'}:
                    conn.execute('update processing_documents set status=?, final_state=?, final_document_id=?, quarantine_reason=? where id=?', (final_state, final_state, final_document_id, reason, document_id))
                file_doc_rows = conn.execute('select status, final_state, quarantine_reason, promotion_blocked_reason from processing_documents where file_id=? order by coalesce(segment_index, 1) asc, id asc', (file_id,)).fetchall()
                states = [str(row['final_state'] or row['status'] or '') for row in file_doc_rows]
                aggregate_error = next((str(row['quarantine_reason'] or row['promotion_blocked_reason'] or '').strip() for row in file_doc_rows if str(row['quarantine_reason'] or row['promotion_blocked_reason'] or '').strip()), '')
                if any(state == 'processing' for state in states):
                    file_status = 'processing'
                elif any(state == 'retry_pending' for state in states):
                    file_status = 'retry_pending'
                elif states and all(state == 'final' for state in states):
                    file_status = 'final'
                    aggregate_error = ''
                elif any(state == 'quarantine' for state in states):
                    file_status = 'quarantine'
                elif any(state == 'unrecognized' for state in states):
                    file_status = 'unrecognized'
                else:
                    file_status = final_state
                conn.execute('update files set status=?, last_error=? where id=?', (file_status, aggregate_error, file_id))
                self._sync_working_document_fts(conn, file_id)
                conn.execute('''insert into document_results (file_id, document_id, final_state, reason, decided_at, correlation_id, final_document_id) values (?, ?, ?, ?, ?, ?, ?) on conflict(file_id) do update set document_id=excluded.document_id, final_state=excluded.final_state, reason=excluded.reason, decided_at=excluded.decided_at, correlation_id=excluded.correlation_id, final_document_id=excluded.final_document_id''', (file_id, document_id, final_state, reason_to_store, utc_now(), correlation_id, final_document_id))
                self._set_queue_size(conn)
                conn.commit()
            except Exception:
                conn.rollback()
                raise
            finally:
                conn.execute('detach database prod_db')
        if output_directory:
            try:
                output_bucket = 'final' if final_state == 'final' else final_state
                self.export_processing_document_to_output(file_id, document_id, Path(output_directory), bucket=output_bucket)
            except Exception:
                self.log_runtime(f'WARN output-copy failed file_id={file_id} state={final_state} output={output_directory}')
        if final_state == 'final':
            self.record_promotion_audit(file_id, document_id, final_document_id, promotion_key, audit_outcome or 'promoted', reason_to_store, correlation_id, {'final_document_id': final_document_id})
        self.set_last_success(f'{final_state.upper()} file_id={file_id}')
        self.log_event('info', 'document.finalized', f'Dokument byl ukoncen stavem {final_state}.', correlation_id, file_id=file_id, document_id=document_id, payload={'reason': reason_to_store})

    def log_runtime(self, message: str) -> None:
        self.runtime_log_path.parent.mkdir(parents=True, exist_ok=True)
        with self.runtime_log_path.open('a', encoding='utf-8') as handle:
            handle.write(f'[{utc_now()}] {message}\n')

    def log_event(self, level: str, event_type: str, message: str, correlation_id: str | None, *, file_id: int | None = None, document_id: int | None = None, attempt_id: int | None = None, payload: dict[str, Any] | None = None) -> None:
        safe_payload = dict(payload or {})
        entry = {
            'schema_version': 2,
            'timestamp': utc_now(),
            'level': level,
            'event_type': event_type,
            'event_name': event_type,
            'message': message,
            'project_path': str(self.project_path),
            'correlation_id': correlation_id,
            'file_id': file_id,
            'document_id': document_id,
            'attempt_id': attempt_id,
            'payload': safe_payload,
        }
        for key in ('endpoint', 'model', 'request_fingerprint', 'duration_ms', 'status_code', 'phase', 'result_code', 'source_sha256'):
            value = safe_payload.get(key)
            if value not in (None, ''):
                entry[key] = value
        self.structured_log_path.parent.mkdir(parents=True, exist_ok=True)
        with self.structured_log_path.open('a', encoding='utf-8') as handle:
            handle.write(json.dumps(entry, ensure_ascii=False) + '\n')
        with self.work_connection() as conn:
            conn.execute('''insert into operational_log(level, event_type, message, correlation_id, file_id, document_id, attempt_id, payload_json) values (?, ?, ?, ?, ?, ?, ?, ?)''', (level, event_type, message, correlation_id, file_id, document_id, attempt_id, json.dumps(safe_payload, ensure_ascii=False)))
            conn.commit()

    def _set_queue_size(self, conn: sqlite3.Connection) -> None:
        queue_size = conn.execute("select count(*) from files where status in ('new','queued','processing','retry_pending')").fetchone()[0]
        conn.execute("insert into system_state(key, value, updated_at) values('queue_size', ?, ?) on conflict(key) do update set value=excluded.value, updated_at=excluded.updated_at", (str(queue_size), utc_now()))

apply_reporting_view_bindings(ProjectRepository)
